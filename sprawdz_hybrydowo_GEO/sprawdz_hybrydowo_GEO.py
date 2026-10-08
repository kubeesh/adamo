# -*- coding: utf-8 -*-
"""
sprawdz_hybrydowo_GEO.py v5 (v4 + kontur DXF vs rysunek + nakładka 1:1) — jak v3 (geometria DXF), PLUS kontrola geometrii 3D z pliku STP
(grubość i gabaryt każdej części z rzeczywistej bryły, nie tylko z opisu/nazwy pliku) ORAZ
czytelny, strukturalny raport zbiorczy .xlsx (kolorowane błędy/uwagi, sortowanie po pozycji BOM
= dokładne miejsce błędu w strukturze łyżki), zamiast samej ściany tekstu w konsoli.

Porównuje krzyżowo: BOM (xlsx) <-> nazwy plików DXF/PDF <-> zawartość PDF
(tabliczka + stempel "Ilość:" z PdfRenamera) <-> geometria DXF (2D, rozwinięcie/kontur)
<-> geometria STP (3D, rzeczywista bryła każdej części w zespole, jeśli plik STP istnieje).

Reguły wyprowadzone z CAŁEGO archiwum S:\\14_Gotowe_Projekty\\GJERSTAD
(154 wydania hybrydowe, ~4900 pozycji BOM; 2026-07-03), plus:
2026-07-18 — dodano kontrolę STP i raport xlsx na prośbę użytkownika ("połącz sprawdzenie
STP-BOM ze sprawdz hybrydowo GEO, żeby wszystkie błędy były w jednym raporcie i były
czytelniejsze / pokazywały gdzie jest błąd").
2026-09-30 — różnice z pliku .stp nie są już BŁĘDAMI, tylko osobnym poziomem "INFO STP" (niebieskie
wiersze): gabaryt bryły z STP zależy od ułożenia części w złożeniu (obrót, pochylenie, gięcie), więc
na GE98022-S2 dał 6 "błędów" przy w 100% poprawnej dokumentacji (każdy odtworzony samym obrotem
płaskiego konturu z DXF, błąd <= 0,7 mm). O zgodności decydują DXF, PDF i BOM. Na końcu arkusza
raportu jest podsumowanie: czy wszystkie DXF są zgodne z BOM + objaśnienie wierszy STP.
2026-10-08 (v5) — kontrola KONTURU DXF vs RYSUNEK PDF (zgłoszenie: "DXF nie pokrywa się z rysunkiem,
fazy zepsute" — v4 porównywała tylko gabaryt z BOM, więc zła faza przechodziła): przerwy i wolne końce
konturu, rozgałęzienia (nieprzycięty narożnik po fazie, linie fazy/ukosu krawędzi), podwójne linie,
elementy dorysowane na warstwie 0, fazy narożników vs notki faz na rysunku ("4x 10 x 45°"), otwory vs
Ø na rysunku, współśrodkowe okręgi (pogłębienie otworu w DXF), gabaryt DXF vs tabliczka PDF, gdy BOM
nie ma wymiarów. Szczegóły i progi: sekcja "kontur DXF vs rysunek" niżej. Testy: testy/test_kontur.py.
2026-10-08 — NAKŁADKA DXF 1:1 na widok z rysunku PDF (wektorowo, bez skanowania obrazu): każdy punkt
konturu DXF musi leżeć na linii rysunku (tol. 1 mm), plus podgląd NAKLADKI_<data>.pdf. Sprawdzone na
prawdziwej części PG4136348 (GE21745AF): zgodna 0,43 mm; zepsucie fazy / przesunięcie wycięcia wykryte.
Testy: testy/test_100_czesci.py (100 części, 14 typów wad, 8 odmian rysunku/DXF).
2026-10-08 (uodpornienie) — ZASADA BEZPIECZEŃSTWA: każdy DXF ma status ZGODNY 1:1 / RÓŻNICE / NIEZWERYFIKOWANY
(z powodem; zakładka "Weryfikacja DXF 1 do 1"), folder "ZGODNY" tylko gdy wszystkie DXF potwierdzone; AUTOTEST
przy starcie; błąd programu przy pozycji/folderze nie przerywa reszty i trafia do raportu jako BŁĄD; szukanie
widoku przez korelację obrazu strony (FFT) gdy skupiska zawiodą; dowolny kąt widoku; tolerancja dopasowana do
dokładności rysunku; bloki INSERT, OCS okręgów, jednostki DXF; notki faz tylko przy widoku rozwinięcia.

  Kategorie części (z DESCRIPTION + flag + struktury BOM):
  - blacha "PL t x a x b" (Normal)          -> DXF + PDF (99% archiwum),
      ALE z flagą C (cięta z płaskownika/lemiesza) -> tylko PDF (76%), DXF dopuszczalny;
      krawędź typu FRONT/CUTTING EDGE z wąskiego paska (szer. <=400) bez DXF -> uwaga.
  - "PL t" / przygotówka PRZ                -> DXF wymagany, PDF opcjonalny (48/47%).
  - lemiesz LM / pręt PR / rura RU/RØR / Round Bar / stock BL -> tylko PDF; DXF dopuszczalny
      (lemiesze HALF ARROW L/P mają DXF, środkowe nie).
  - złożenie (flaga M)                      -> PDF (__X__UNKNOWN__m); DXF przy złożeniu = uwaga
      (loadery G BOLTON mają DXF złożenia - szablon).
  - ramka tabliczki (DEK-*/PG4104777, title FRAME..SIGN) -> PDF bywa albo nie; tylko uwaga.
  - Purchased / STICKER                     -> pliki opcjonalne (93% bez plików).

  Zgodność metadanych:
  - grubość w nazwie musi być JEDNYM z wymiarów DESCRIPTION (kolejność t/a/b bywa różna,
    np. "PL 123 x 90 x 3" = grubość 3; "PL 60 x 186 x 35" = grubość 60);
  - ilość w nazwie/stemplu = QTY wiersza LUB ilość całkowita (QTY x rodzice wg POS,
    np. poz. 7.3.2 przy rodzicu x2 -> pliki nazwane 2 przy QTY=1);
  - flagi w nazwie = kolumny M/O/P/U/C (m=montaż, o=obróbka, p=gięcie, u=ukosowanie, c=cięcie);
  - stempel "Ilość: N" != QTY jest OK przy adnotacji o odbiciu lustrzanym
    ("DRUGA SZTUKA ODBICIE LUSTRZANE");
  - DXF ma prefiks głównego złożenia (numer z BOM, NIE nazwa folderu), PDF bez prefiksu;
  - pliki "P10.01 WYTYCZNE GJERSTAD" itp. ignorowane.

  Geometria STP (opcjonalna, wymaga Inventora zainstalowanego lokalnie + pywin32):
  - szuka pliku "<nazwa_folderu_bez_dopiskow>.stp"/".step" w folderze projektu, albo
    dowolnego *.stp/*.step jeśli nazwa nie pasuje (znaleziono 182 takie pliki w archiwum
    GJERSTAD 2026-07-18, w większości nazwane jak folder projektu, np. GE6725\\GE6725.stp);
  - otwiera plik przez COM (win32com) w Inventorze (używa już uruchomionej instancji, jeśli
    jest, inaczej odpala nową w tle); pliki STP z tego archiwum zachowują strukturę zespołu
    (PRODUCT('PG4104413','PG4104413','PL 12 x 1633 x 1406',...) w danych STEP) — Inventor
    powinien zaimportować to jako ZESPÓŁ z osobnymi komponentami, nie jedną zlaną bryłę;
  - dla każdej liściastej części czyta LOKALNY (nieobrócony) RangeBox definicji części
    (PartComponentDefinition.RangeBox — w układzie współrzędnych samej części, NIE zespołu,
    więc obrót/pozycja części w złożeniu nie zniekształca wyniku, w przeciwieństwie do
    RangeBox samego wystąpienia w przestrzeni zespołu);
  - 3 wymiary lokalnego pudełka porównywane wprost (posortowane) z 3 wymiarami z DESCRIPTION
    BOM, tą samą tolerancją co geometria DXF (GEO_TOL_BLAD/GEO_TOL_INFO), z tym samym
    złagodzeniem dla części giętych (flaga P) i obrabianych (flagi O/U) — patrz komentarz
    przy geo_porownaj_3d.
  - ZWERYFIKOWANE NA ŻYWO 2026-07-18 (COM przez Inventor 2026.1, plik GE75206-S1.stp z tego
    archiwum): Inventor faktycznie importuje ten STEP jako ZESPÓŁ (DocumentType=12291) z 64
    osobnymi wystąpieniami, iProperty "Part Number" poprawnie = numer części z BOM (np.
    'PG4104333'). Jedyna pułapka: `Documents.Open()` zwraca ogólny obiekt "Document" bez
    `.ComponentDefinition` (early-bound gen_py nie widzi interfejsu AssemblyDocument) —
    rozwiązanie: `win32com.client.CastTo(oDoc, "AssemblyDocument")` (oficjalny mechanizm
    pywin32 na ten dokładnie przypadek). Sprawdzone wyniki dla GE75206-S1 są SPÓJNE z
    niezależnie znalezionym błędem tabliczki PDF: bryła STP `PG4104333` (BOTTOM PLATE) =
    1226.0 mm, dokładnie zgadza się z tabliczką PDF (1226), a NIE z BOM (1230) — czyli to
    naprawdę BOM ma nieaktualną wartość, nie model/PDF. Części gięte (flaga P) mają lokalny
    RangeBox mniejszy niż DESCRIPTION (bo bryła jest już wygięta, a DESCRIPTION to zwykle
    rozwinięcie) - stąd złagodzenie do UWAGI dla flagi P/O/U, tak samo jak przy DXF.

Użycie: sprawdz_hybrydowo_GEO.py <folder> [folder2 ...]  (albo przeciągnij folder(y) na .bat)
Wymaga: pip install openpyxl pymupdf ezdxf pywin32
Na końcu zapisuje jeden zbiorczy raport .xlsx (obok pierwszego podanego folderu) i otwiera go.
"""
import atexit
import datetime
import math
import os
import re
import sys

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    import fitz  # pymupdf
except ImportError as e:
    print("Brak biblioteki:", e, "->  pip install openpyxl pymupdf")
    sys.exit(2)

try:
    import ezdxf
    import numpy as np  # zależność ezdxf, więc jest zawsze razem z nim
    from ezdxf import bbox as ezbbox
    from ezdxf import path as ezpath
    GEO = True
except ImportError:
    GEO = False  # bez ezdxf dziala jak zwykla wersja v2 (bez kontroli geometrii DXF)

try:
    import win32com.client
    STP_MODULE_OK = True
except ImportError:
    STP_MODULE_OK = False  # bez pywin32 dziala bez kontroli geometrii STP

# tolerancje kontroli geometrii (bbox konturu DXF/bryly STP vs wymiary DESCRIPTION)
GEO_TOL_BLAD = 3.0   # mm — odchylka powyzej -> BLAD
GEO_TOL_INFO = 1.5   # mm — strefa szara -> UWAGA

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

FLAG_LETTERS = set("mopuc")
JUNK = re.compile(r"WYTYCZNE|^__?P10\.01", re.I)

# Aliasy nazw kolumn BOM w roznych wariantach jezykowych/eksportow Inventora - zbudowane
# 2026-07-19 ze skanu 38 realnych plikow xlsx w S:\14_Gotowe_Projekty\GJERSTAD (12 wariantow
# naglowka: angielski ITEM/PART NUMBER/QTY/..., polski POZYCJA/NUMER CZĘŚCI/ILOŚĆ/...).
# Kolumna "TEMAT" (widziana obok TYTUŁ w wariancie polskim) to INNE pole niz TITLE
# (najprawdopodobniej "Subject") - celowo pominieta jako alias.
ALIASY_KOLUMN = {
    "POS": ("POS", "ITEM", "POZYCJA"),
    "PART": ("PART NUMBER", "NUMER CZĘŚCI", "NUMER CZESCI"),
    "STRUCT": ("BOM STRUCTURE", "TYP W STRUKTURZE ZESTAWIENIA BOM"),
    "QTY": ("QTY", "ILOŚĆ", "ILOSC"),
    "DESC": ("DESCRIPTION", "OPIS"),
    "MAT": ("MATERIAL", "MATERIAŁ"),
    "TITLE": ("TITLE", "TYTUŁ", "TYTUL"),
}

# zbiorczy raport strukturalny (wszystkie foldery podane w jednym uruchomieniu)
WSZYSTKIE_WIERSZE = []   # list of dict: folder, pos, part, tytul, poziom, plik, rodzaj, opis
PODSUMOWANIE_FOLDEROW = []  # list of dict: folder, ok, bledy, uwagi
WERYFIKACJA_DXF = []     # list of dict: jeden wpis na każdy DXF części — status 1:1 z rysunkiem i dowód

# ZASADA BEZPIECZEŃSTWA (2026-10-08): DXF jest "ZGODNY 1:1" tylko wtedy, gdy program POZYTYWNIE potwierdził
# go nakładką na rysunek i kontur jest czysty. Wszystko inne (brak PDF, brak widoku, błąd programu,
# odbicie lustrzane, obróbka z naddatkiem) to "NIEZWERYFIKOWANY" z powodem — nigdy cichy sukces.
AUTOTEST = dict(ok=True, opis="nie uruchamiany (moduł zaimportowany)")
PODGLAD_WSZYSTKIE = os.environ.get("SPRAWDZ_PODGLAD_WSZYSTKIE") == "1"          # podgląd nakładki każdej części
NIEZWERYFIKOWANE_BLOKUJA = os.environ.get("SPRAWDZ_NIEZWERYFIKOWANE_BLAD") == "1"  # kod wyjścia 1 przy niezweryf.
BEZ_STP = os.environ.get("SPRAWDZ_BEZ_STP") == "1"   # pomiń model .stp (bez uruchamiania Inventora)


# ---------------------------------------------------------------- BOM
def desc_dims(desc):
    """Zbior wymiarow z DESCRIPTION 'PL t x a x b' (kolejnosc bywa dowolna)."""
    m = re.match(r"PL\.?\s*([\d,\.]+(?:\s*[xX]\s*[\d,\.]+)*)", desc.strip())
    if not m:
        return []
    dims = []
    for x in re.split(r"\s*[xX]\s*", m.group(1)):
        x = x.replace(",", ".").strip(".")
        try:
            dims.append(float(x))
        except ValueError:
            pass
    return dims


def _znajdz_kolumne_m(hdr):
    """Kolumna startowa flag M/O/P/U/C. Dopasowuje 'M' oraz warianty z dopiskiem typu
    'M\\n(MONTAŻ)' (znaleziono w archiwum), ale NIE 'MATERIAL'/'MATERIAŁ'/'MINIATURA'."""
    for i, h in enumerate(hdr):
        if re.match(r'^M(\s|\(|$)', h.strip()):
            return i
    return None


def czytaj_bom(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rows = []
    for r in ws.iter_rows(values_only=True):
        rows.append(list(r))
        if len(rows) > 500:
            break
    wb.close()
    hdr_i = None
    for i, r in enumerate(rows[:6]):
        up = [str(c or "").strip().upper() for c in r]
        if (any(a in up for a in ALIASY_KOLUMN["PART"])
                and any(a in up for a in ALIASY_KOLUMN["QTY"])):
            hdr_i = i
            break
    if hdr_i is None:
        return None, "nie znaleziono nagłówka BOM (PART NUMBER/QTY, w żadnym znanym wariancie językowym)"
    hdr = [str(c or "").strip().upper() for c in rows[hdr_i]]

    def col(field, default=None):
        for alias in ALIASY_KOLUMN[field]:
            if alias in hdr:
                return hdr.index(alias)
        return default

    c_pos = col("POS", default=0)
    c_part = col("PART", default=1)
    c_struct = col("STRUCT", default=4)
    c_qty = col("QTY", default=5)
    c_desc = col("DESC", default=7)
    c_mat = col("MAT", default=10)
    c_title = col("TITLE")
    c_m = _znajdz_kolumne_m(hdr)
    if c_m is None:
        c_m = 11  # fallback jak w wersji bez rozpoznanych aliasow
    bom = []
    for r in rows[hdr_i + 1:]:
        if c_part >= len(r) or not r[c_part]:
            continue
        desc = str(r[c_desc] or "").strip() if c_desc < len(r) else ""
        bom.append(dict(
            pos=str(r[c_pos]).strip() if r[c_pos] is not None else "?",
            part=str(r[c_part]).strip(),
            struct=str(r[c_struct] or "").strip() if c_struct < len(r) else "",
            qty=r[c_qty] if c_qty < len(r) else None,
            desc=desc, dims=desc_dims(desc),
            mat=str(r[c_mat] or "").strip() if c_mat < len(r) else "",
            title=str(r[c_title] or "").strip() if (c_title is not None and c_title < len(r)) else "",
            flags={f for off, f in enumerate("mopuc") if c_m + off < len(r) and r[c_m + off]},
        ))
    # ilosc calkowita: iloczyn QTY rodzicow wg POS (poz. "7.3.2" -> rodzice "7.3", "7")
    qty_by_pos = {b["pos"]: b["qty"] for b in bom}
    for b in bom:
        total = b["qty"] if isinstance(b["qty"], (int, float)) else None
        if total is not None and "." in b["pos"]:
            czesci = b["pos"].split(".")
            for k in range(1, len(czesci)):
                q = qty_by_pos.get(".".join(czesci[:k]))
                if isinstance(q, (int, float)):
                    total *= q
        b["qty_total"] = total
    return bom, None


def kategoria(b, proj_root):
    d = b["desc"].upper()
    t = (b["title"] or "").upper()
    part = b["part"].upper()
    if b["struct"].lower() == "purchased" or "STICKER" in part or "STICKER" in t:
        return "purchased"
    if part.startswith("DEK-") and "FRAME" in t or ("FRAME" in t and "SIGN" in t) or part == "PG4104777":
        return "ramka"
    if "m" in b["flags"] or part == proj_root:
        return "zlozenie"
    if d.startswith("LM") or re.match(r"(PR\d|PR\s|RU\d|RUND|ROUND\s*BAR|RØR|ROR\b|AKSEL)", d) or d.startswith("BL"):
        return "material_ciety"      # lemiesz/pret/rura/stock: PDF, DXF opcjonalny
    if re.match(r"PL\.?\s*[\d,\.]+\s*[xX]\s*[\d,\.]+\s*[xX]", d):
        s = sorted(b["dims"])
        if len(s) == 3 and s[0] == s[1] and s[2] >= 3 * s[0]:
            return "material_ciety"  # przekroj kwadratowy (np. PL 80 x 80 x 692) = pret, nie blacha
        if "c" in b["flags"]:
            return "blacha_c"        # cieta z plaskownika: PDF, DXF opcjonalny
        return "blacha"              # DXF + PDF
    if re.match(r"PL\.?\s*[\d,\.]+(\s*[xX]\s*[\d,\.]+)?\s*(MM)?\s*$", d):
        return "przygotowka"         # DXF, PDF opcjonalny
    if "c" in b["flags"]:
        return "material_ciety"
    return "inne"                    # PDF oczekiwany


def czy_krawedz_plaskownik(b):
    """FRONT/CUTTING EDGE z waskiego paska (szer<=400): brak DXF czesty w archiwum (grading)."""
    t = (b["title"] or "").upper()
    if not re.search(r"FRONT\s*EDGE|CUTTING\s*EDGE|KNIFE", t):
        return False
    dims = sorted(b["dims"], reverse=True)
    return len(dims) == 3 and dims[1] <= 400


# ---------------------------------------------------------------- nazwy plikow
def parsuj_nazwe(fname, bom_parts, bom_asm):
    base, ext = os.path.splitext(fname)
    ext = ext.upper()
    toks = base.split("__")
    while toks and toks[0] == "":
        toks.pop(0)
    if not toks:
        return None

    def meta_like(t):
        return (t == "" or t.upper() == "UNKNOWN" or t.isdigit()
                or re.fullmatch(r"[\d,\.]+\s*mm", t)
                or (len(t) == 1 and t.lower() in FLAG_LETTERS))

    pi = 0
    if len(toks) > 1 and not meta_like(toks[1]):
        t1u = toks[1].upper()
        if (t1u in bom_parts
                or (t1u.endswith("PRZ") and t1u[:-3] in bom_parts)
                or toks[0].upper() in bom_asm):
            pi = 1
    part, rest = toks[pi], toks[pi + 1:]
    prefiks = pi == 1
    if rest and rest[0] == "UNKNOWN":
        return dict(part=part, thk=None, mat=None, qty=None,
                    flags={t.lower() for t in rest[1:] if t}, asm=True, prefiks=prefiks, ext=ext)
    thk = mat = qty = None
    flags = set()
    for t in rest:
        if not t:
            continue
        m = re.fullmatch(r"([\d,\.]+)\s*mm", t)
        if m and thk is None:
            thk = float(m.group(1).replace(",", "."))
        elif len(t) == 1 and t.lower() in FLAG_LETTERS:
            flags.add(t.lower())
        elif t.isdigit():
            qty = int(t)      # ostatni czysto-liczbowy token = ilosc
        elif thk is None and re.fullmatch(r"\d+[,\.]\d+", t):
            thk = float(t.replace(",", "."))  # grubosc zapisana bez "mm" (starsze wydania)
        elif mat is None:
            mat = t
    return dict(part=part, thk=thk, mat=mat, qty=qty, flags=flags,
                asm=False, prefiks=prefiks, ext=ext)


def mat_zgodny(a, b):
    if not a or not b:
        return True
    def norm(x):
        x = re.sub(r"[\s\.]", "", x).upper()
        return x.replace("HARDOX", "HB").replace("HX", "HB")  # Hardox 500 == HB500 == HX500
    na, nb = norm(a), norm(b)
    return na.startswith(nb) or nb.startswith(na)


# ---------------------------------------------------------------- PDF
def czytaj_pdf(path):
    doc = fitz.open(path)
    text = "\n".join(doc[i].get_text() for i in range(len(doc)))
    notki = []          # notki faz z położeniem: (strona, x, y, nogi) — do odróżnienia faz narożnika od krawędzi
    for nr, strona in enumerate(doc):
        for blok in strona.get_text("dict")["blocks"]:
            for linia in blok.get("lines", []):
                t = "".join(sp["text"] for sp in linia["spans"])
                for f in wymiary_rysunku(t)["fazy"]:
                    x0, y0, x1, y1 = linia["bbox"]
                    notki.append(dict(f, strona=nr, x=(x0 + x1) / 2, y=(y0 + y1) / 2))
    doc.close()
    out = {}
    m = re.search(r"Ilość\s*:\s*(\d+)", text)
    if m:
        out["stamp"] = int(m.group(1))
    if re.search(r"ODBICI|LUSTRZ|MIRROR", text, re.I):
        out["lustro"] = True
    m = re.search(r"PL\.?\s*([\d,\.]+\s*x\s*[\d,\.]+\s*x\s*[\d,\.]+)", text)
    if m:
        dims = []
        for x in re.split(r"\s*x\s*", m.group(1)):
            try:
                dims.append(float(x.replace(",", ".").strip(".")))
            except ValueError:
                pass
        if len(dims) == 3:
            out["dims"] = sorted(dims)
    m = re.search(r"(S355(?:J2G3|J2|MC)?|HARDOX\s*\d{3}|HB\s*\d{3})", text, re.I)
    if m:
        out["mat"] = m.group(1)
    out.update(wymiary_rysunku(text))
    out["notki_faz"] = notki
    return out


LICZBA = r"(\d+(?:[.,]\d+)?)"


def _liczba(s):
    return float(s.replace(",", "."))


def wymiary_rysunku(text):
    """Wymiary z tekstu rysunku PDF potrzebne do porównania z DXF:
    - fazy: notki Inventora "5 x 45°", "2x 5 x 45°" (ilość z przodu), "FAZA 5x5";
      nogi fazy = (d, d*tg(kąt)), posortowane;
    - fi: średnice otworów "Ø10,5" — Inventor pisze Ø czcionką AIGDT, która w tekście PDF
      daje literę "n" ("n10,5" = Ø10,5);
    - gwinty: "M12" (w DXF jest wtedy otwór pod gwint, ~0,75-1,0 x M);
    - promienie: "R15" (duże wycięcia okrągłe bywają wymiarowane promieniem);
    - liczby: wszystkie liczby z rysunku (do sprawdzenia, czy faza z DXF jest zwymiarowana)."""
    fazy = []
    for m in re.finditer(r"(?<![\d.,])(?:(\d+)\s*[xX×]\s*)?" + LICZBA + r"\s*(?:mm)?\s*[xX×]\s*"
                         + LICZBA + r"\s*(?:deg|[°º˚])", text):
        d, kat = _liczba(m.group(2)), _liczba(m.group(3))
        if d <= 0 or not 0 < kat < 90:
            continue
        nogi = tuple(sorted((round(d, 2), round(d * math.tan(math.radians(kat)), 2))))
        fazy.append(dict(nogi=nogi, ile=int(m.group(1)) if m.group(1) else None,
                         tekst=re.sub(r"\s+", " ", m.group(0)).strip()))
    for m in re.finditer(r"FAZ[AY]?\s+" + LICZBA + r"\s*[xX×]\s*" + LICZBA
                         + r"(?![\d.,])(?!\s*(?:deg|[°º˚]))", text, re.I):
        a, b = _liczba(m.group(1)), _liczba(m.group(2))
        nogi = (a, a) if b == 45 else tuple(sorted((a, b)))
        fazy.append(dict(nogi=nogi, ile=None, tekst=re.sub(r"\s+", " ", m.group(0)).strip()))
    fi = [_liczba(x) for x in re.findall(r"(?:[Ø⌀ø∅]|(?:(?<=[\s\dxX×(])|^)n)\s?" + LICZBA, text, re.M)]
    gwinty = [_liczba(x) for x in re.findall(r"(?<![A-Za-z])M" + LICZBA, text)]
    promienie = [_liczba(x) for x in re.findall(r"(?<![A-Za-z])R\s?" + LICZBA, text)]
    liczby = {_liczba(x) for x in re.findall(r"\d+(?:[.,]\d+)?", text)}
    return dict(fazy=fazy, fi=fi, gwinty=gwinty, promienie=promienie, liczby=liczby)


# ---------------------------------------------------------------- geometria DXF
def dxf_gabaryt(path):
    """Kandydackie obrysy konturu DXF: bbox osiowy + prostokąty we wszystkich
    orientacjach krawędzi otoczki wypukłej. DESCRIPTION Inventora to gabaryt
    w osiach części — przy obróconym eksporcie (SIDE KNIFE) albo części
    trójkątnej (SIDE MEMBER) właściwa orientacja to któraś z krawędzi konturu."""
    # 2026-10-08: tylko warstwy cięcia (bez wymiarów, ramek, linii ukrytych/gięcia) i z rozwiniętymi blokami —
    # te same elementy, które zobaczy wypalarka i które sprawdza nakładka/kontur
    pts = []
    for p in _prymitywy_dxf(ezdxf.readfile(path)):
        krok = max(1, len(p["pts"]) // 60)
        pts.extend((round(float(x), 2), round(float(y), 2)) for x, y in np.vstack([p["pts"][::krok], p["pts"][-1:]]))
    if not pts:
        return []
    xs = [p[0] for p in pts]; ys = [p[1] for p in pts]
    kand = [(max(xs) - min(xs), max(ys) - min(ys))]
    spts = sorted(set(pts))
    if len(spts) <= 2:
        return kand
    def cr(o, a, b):
        return (a[0]-o[0])*(b[1]-o[1]) - (a[1]-o[1])*(b[0]-o[0])
    lo = []
    for p in spts:
        while len(lo) >= 2 and cr(lo[-2], lo[-1], p) <= 0:
            lo.pop()
        lo.append(p)
    up = []
    for p in reversed(spts):
        while len(up) >= 2 and cr(up[-2], up[-1], p) <= 0:
            up.pop()
        up.append(p)
    hull = lo[:-1] + up[:-1]
    n = len(hull)
    for i in range(n):
        x1, y1 = hull[i]; x2, y2 = hull[(i + 1) % n]
        dx, dy = x2 - x1, y2 - y1
        L = math.hypot(dx, dy)
        if L < 1e-9:
            continue
        ux, uy = dx / L, dy / L
        us = [p[0]*ux + p[1]*uy for p in hull]
        vs = [-p[0]*uy + p[1]*ux for p in hull]
        kand.append((max(us) - min(us), max(vs) - min(vs)))
    return kand


def geo_porownaj(b, dxf_info, kand):
    """Najlepsze dopasowanie kandydatów obrysu DXF do wymiarów płaskich z DESCRIPTION.
    Zwraca (odchylka, plan_dims, obrys)."""
    dims = list(b["dims"])
    warianty = []
    if dxf_info["thk"] is not None and dxf_info["thk"] in dims:
        d2 = list(dims)
        d2.remove(dxf_info["thk"])             # grubość znana z nazwy pliku
        warianty.append(sorted(d2))
    else:
        for i in range(3):
            warianty.append(sorted(dims[:i] + dims[i + 1:]))
    best = None
    for plan in warianty:
        for (w, h) in kand:
            bbs = sorted((w, h))
            dev = max(abs(plan[0] - bbs[0]), abs(plan[1] - bbs[1]))
            if best is None or dev < best[0]:
                best = (dev, plan, bbs)
    return best


def geo_porownaj_3d(b_dims, dims3):
    """Jak geo_porownaj, ale dla bryły 3D z STP: mamy already 3 prawdziwe wymiary lokalnego
    pudełka (nie trzeba zgadywać orientacji jak przy płaskim DXF) - proste porównanie
    posortowanych trójek. Zwraca (odchylka, plan_posortowany, znaleziony_posortowany)
    albo None jeśli brak 3 wymiarów po którejkolwiek stronie."""
    plan = sorted(b_dims)
    znaleziony = sorted(dims3)
    if len(plan) != 3 or len(znaleziony) != 3:
        return None
    dev = max(abs(p - f) for p, f in zip(plan, znaleziony))
    return dev, plan, znaleziony


# ---------------------------------------------------------------- kontur DXF vs rysunek (fazy, otwory)
# 2026-10-08 — przełożony zgłosił DXF niezgodny z rysunkiem i "zepsute fazy", których program nie
# wyłapał: wcześniejsza kontrola porównywała tylko GABARYT konturu z BOM, a faza 5 mm prawie nie
# zmienia gabarytu. Ta część czyta sam kontur, tak jak zobaczy go wypalarka:
#   - przerwy / wolne końce linii (kontur niedomknięty, linia wystaje za narożnik),
#   - rozgałęzienia (linia kończy się na środku innej albo 3+ linie w jednym punkcie) — typowy
#     ślad fazy dorysowanej bez przycięcia starego narożnika, albo linii fazy/ukosu krawędzi
#     wyeksportowanej z widoku, którą laser by wyciął,
#   - podwójne linie, elementy dorysowane ręcznie na warstwie "0" (opisane w PORADNIKU od
#     03.07.2026, ale w wersji v4 nie było ich w kodzie),
#   - fazy narożników i otwory z DXF porównane z notkami/wymiarami na rysunku PDF,
#   - współśrodkowe okręgi (faza/pogłębienie otworu wyeksportowane jako drugi okrąg).
STYK_TOL = 0.05      # mm — końce bliżej siebie = połączone (CAM i tak je skleja)
PRZERWA_MAX = 2.0    # mm — dwa wolne końce bliżej siebie = "przerwa", dalej = "wolny koniec"
FAZA_MAX = 60.0      # mm — dłuższy odcinek nie jest traktowany jako faza narożnika
FAZA_TOL = 0.5       # mm — tolerancja nogi fazy DXF vs rysunek
FAZA_MIN_RYS = 2.0   # mm — fazy z rysunku do tej wielkości to łamanie krawędzi, nie kontur
OTWOR_TOL = 0.15     # mm — tolerancja średnicy otworu DXF vs Ø z rysunku
WYSTAJE_MAX = 15.0   # mm — dorysowany element dalej poza obrysem = zbłąkana geometria

# warstwy, które nie są cięte (gięcie, linie ukryte/osie/styczne, wymiary, ramki, trasowanie);
# warstwa "Widoczne wąskie" to m.in. 3/4 okręgu gwintu — otwarty łuk, nie kontur
WARSTWY_NIE_TNACE = re.compile(
    r"BEND|TANGENT|TOOL_CENTER|ARC_CENTER|ALTREP|UNCONSUMED|ROLL|FEATURE|GI[EĘ]CI|"
    r"UKRYT|HIDDEN|OSI|O[SŚ]\b|CENTER|CENTRE|SYMETR|STYCZN|W[AĄ]SK|"
    r"WYMIAR|DIM|TEXT|TEKST|ANNOT|TYTU|TITLE|RAMK|BORDER|GRANIC|KRESK|HATCH|DEFPOINTS|"
    r"TRAS|GRAW|GRAV|NAPIS|MARK", re.I)


def _klastry(punkty, tol):
    """Grupuje punkty leżące bliżej niż tol (siatka + union-find). Zwraca listę list indeksów."""
    rodzic = list(range(len(punkty)))

    def korzen(i):
        while rodzic[i] != i:
            rodzic[i] = rodzic[rodzic[i]]
            i = rodzic[i]
        return i

    siatka = {}
    for i, (x, y) in enumerate(punkty):
        siatka.setdefault((math.floor(x / tol), math.floor(y / tol)), []).append(i)
    for (cx, cy), idx in siatka.items():
        for dx in (-1, 0, 1):
            for dy in (-1, 0, 1):
                for j in siatka.get((cx + dx, cy + dy), ()):
                    for i in idx:
                        if i < j and math.dist(punkty[i], punkty[j]) <= tol:
                            rodzic[korzen(i)] = korzen(j)
    grupy = {}
    for i in range(len(punkty)):
        grupy.setdefault(korzen(i), []).append(i)
    return list(grupy.values())


def _prymitywy_dxf(doc):
    """Elementy tnące z modelspace rozbite na LINE/ARC/CIRCLE/SPLINE/ELLIPSE (polilinie
    rozbijane na odcinki i łuki, bloki INSERT rozwijane — elementy na warstwie "0" w bloku dziedziczą
    warstwę wstawienia, jak w AutoCADzie), każdy jako łamana pts (krok 0,01 mm na łukach)."""
    out = []

    def dodaj(e, warstwa, glebokosc=0):
        t = e.dxftype()
        if t == "INSERT":
            if glebokosc > 8:
                return
            try:
                for v in e.virtual_entities():
                    w = v.dxf.get("layer", "0")
                    w = warstwa if w == "0" else w
                    if not WARSTWY_NIE_TNACE.search(w):
                        dodaj(v, w, glebokosc + 1)
            except Exception:
                pass
            return
        if t in ("LWPOLYLINE", "POLYLINE"):
            try:
                for v in e.virtual_entities():
                    dodaj(v, warstwa)
            except Exception:
                pass
            return
        if t not in ("LINE", "ARC", "CIRCLE", "SPLINE", "ELLIPSE"):
            return
        try:
            pts = np.array([(v.x, v.y) for v in ezpath.make_path(e).flattening(0.01)])
        except Exception:
            return
        if len(pts) < 2:
            return
        dl = float(np.hypot(*np.diff(pts, axis=0).T).sum())
        if dl < STYK_TOL:
            return  # "linia" długości 0 — artefakt eksportu
        p = dict(typ=t, warstwa=warstwa, pts=pts, dl=dl, reczny=warstwa == "0",
                 zamk=t == "CIRCLE" or math.dist(pts[0], pts[-1]) < STYK_TOL)
        if t == "CIRCLE":
            c = e.ocs().to_wcs(e.dxf.center)       # okrąg z wektorem wyciągnięcia (0,0,-1) ma środek w OCS
            p["srodek"] = (c.x, c.y)
            p["fi"] = 2 * e.dxf.radius
        out.append(p)

    for e in doc.modelspace():
        warstwa = e.dxf.get("layer", "0")
        if not WARSTWY_NIE_TNACE.search(warstwa):
            dodaj(e, warstwa)
    return out


def _sygnatura(p):
    """Ten sam element narysowany 2x daje tę samą sygnaturę (niezależnie od kierunku)."""
    if p["typ"] == "CIRCLE":
        return ("C", round(p["srodek"][0], 2), round(p["srodek"][1], 2), round(p["fi"], 2))
    konce = sorted((tuple(np.round(p["pts"][0], 2)), tuple(np.round(p["pts"][-1], 2))))
    return (p["typ"], tuple(konce), round(p["dl"], 1), tuple(np.round(p["pts"].mean(axis=0), 1)))


def _faza(S, sasiad_a, sasiad_b):
    """Odcinek S między dwoma odcinkami = faza narożnika, jeśli przedłużenia sąsiadów
    przecinają się za S (wirtualny narożnik V) pod kątem 45-135°. Zwraca (noga1, noga2, V)."""
    a, b = S["pts"][0], S["pts"][-1]
    (E1, f1), (E2, f2) = sasiad_a, sasiad_b
    u1 = (f1 - a) / np.linalg.norm(f1 - a)
    u2 = (f2 - b) / np.linalg.norm(f2 - b)
    if abs(float(np.dot(u1, u2))) > 0.7072:   # kąt narożnika poza 45-135°
        return None
    try:
        t1, t2 = np.linalg.solve(np.array([-u1, u2]).T, b - a)  # a - t1*u1 = b - t2*u2 = V
    except np.linalg.LinAlgError:
        return None
    if not (0.1 < t1 <= FAZA_MAX and 0.1 < t2 <= FAZA_MAX):
        return None
    if S["dl"] >= min(E1["dl"], E2["dl"]):
        return None
    nogi = sorted((round(float(t1), 1), round(float(t2), 1)))
    return nogi[0], nogi[1], tuple(a - t1 * u1)


JEDNOSTKI_DXF = {1: "calach", 2: "stopach", 5: "centymetrach", 6: "metrach", 14: "decymetrach"}


def dxf_analiza(path):
    """Kontur DXF tak, jak zobaczy go wypalarka. Zwraca słownik z listami problemów
    (punkty w układzie współrzędnych DXF, żeby dało się je znaleźć w AutoCADzie)."""
    doc = ezdxf.readfile(path)
    prym = _prymitywy_dxf(doc)
    wynik = dict(duble=0, przerwy=[], wolne=[], rozgal=[], fazy=[], okregi=[],
                 wspolsrodkowe=[], reczne=0, wystaje=[], jednostki=doc.header.get("$INSUNITS", 0),
                 pusty=not prym)
    if not prym:
        return wynik

    # --- podwójne linie (zdublowane elementy usuwamy przed analizą konturu)
    widziane, unikalne = set(), []
    for p in prym:
        s = _sygnatura(p)
        if s in widziane:
            wynik["duble"] += 1
        else:
            widziane.add(s)
            unikalne.append(p)

    # --- węzły konturu: każdy koniec otwartego elementu powinien stykać się z dokładnie jednym innym
    otwarte = [p for p in unikalne if not p["zamk"]]
    konce = []  # konce[2*i] = początek otwarte[i], konce[2*i+1] = koniec
    for p in otwarte:
        konce.append(tuple(p["pts"][0]))
        konce.append(tuple(p["pts"][-1]))
    grupy = _klastry(konce, STYK_TOL)
    wezel = {}
    for g in grupy:
        for k in g:
            wezel[k] = g

    seg_a = np.vstack([p["pts"][:-1] for p in unikalne])
    seg_b = np.vstack([p["pts"][1:] for p in unikalne])
    seg_wl = np.concatenate([np.full(len(p["pts"]) - 1, i) for i, p in enumerate(unikalne)])
    idx_w_unikalnych = {id(p): i for i, p in enumerate(unikalne)}

    def na_innym_elemencie(pkt, wlasny):
        d = seg_b - seg_a
        dd = (d * d).sum(axis=1)
        dd[dd == 0] = 1e-12
        t = np.clip(((pkt - seg_a) * d).sum(axis=1) / dd, 0, 1)
        odl = np.hypot(*(seg_a + d * t[:, None] - pkt).T)
        return bool(((odl < STYK_TOL) & (seg_wl != wlasny)).any())

    wolne = []
    for g in grupy:
        pkt = np.mean([konce[k] for k in g], axis=0)
        if len(g) >= 3:
            wynik["rozgal"].append(tuple(pkt))
        elif len(g) == 1:
            if na_innym_elemencie(pkt, idx_w_unikalnych[id(otwarte[g[0] // 2])]):
                wynik["rozgal"].append(tuple(pkt))   # koniec linii na środku innej linii (T)
            else:
                wolne.append(tuple(pkt))
    # wolne końce blisko siebie = przerwa w konturze; reszta = linia urwana / wystająca
    uzyte = set()
    for i, p in enumerate(wolne):
        if i in uzyte:
            continue
        najbl = min(((math.dist(p, q), j) for j, q in enumerate(wolne) if j != i and j not in uzyte),
                    default=None)
        if najbl and najbl[0] <= PRZERWA_MAX:
            uzyte.update((i, najbl[1]))
            wynik["przerwy"].append((najbl[0], p))
        else:
            uzyte.add(i)
            wynik["wolne"].append(p)

    # --- fazy narożników. Współliniowe odcinki stykające się końcami (krawędź podzielona przez
    # eksport albo ręcznie) najpierw łączymy w jeden prosty odcinek.
    linia = [p["typ"] == "LINE" for p in otwarte]
    rodzic = list(range(len(otwarte)))

    def korzen(i):
        while rodzic[i] != i:
            rodzic[i] = rodzic[rodzic[i]]
            i = rodzic[i]
        return i

    def kierunek(i):
        d = otwarte[i]["pts"][-1] - otwarte[i]["pts"][0]
        return d / np.linalg.norm(d)

    for g in grupy:
        if len(g) == 2 and linia[g[0] // 2] and linia[g[1] // 2] and g[0] // 2 != g[1] // 2:
            u, v = kierunek(g[0] // 2), kierunek(g[1] // 2)
            if abs(float(u[0] * v[1] - u[1] * v[0])) < 1e-3:   # współliniowe
                rodzic[korzen(g[0] // 2)] = korzen(g[1] // 2)
    proste = {}   # korzeń -> końce zewnętrzne (indeksy w konce) prostego odcinka
    for i in range(len(otwarte)):
        if not linia[i]:
            continue
        for e in (0, 1):
            g = wezel[2 * i + e]
            wew = len(g) == 2 and any(k // 2 != i and linia[k // 2] and korzen(k // 2) == korzen(i) for k in g)
            if not wew:
                proste.setdefault(korzen(i), []).append(2 * i + e)

    def prosty(k):
        """Prosty odcinek (po scaleniu), do którego należy koniec k: (pkt_k, pkt_drugi, długość)."""
        ends = proste.get(korzen(k // 2), [])
        if len(ends) != 2:
            return None
        drugi = ends[1] if ends[0] == k else ends[0]
        a, b = np.array(konce[k]), np.array(konce[drugi])
        return a, b, float(np.linalg.norm(b - a))

    for r, ends in proste.items():
        if len(ends) != 2:
            continue
        a, b, dl = prosty(ends[0])
        if dl > FAZA_MAX:
            continue
        sasiedzi = []
        for k in ends:
            g = wezel[k]
            if len(g) != 2:
                break
            inny = g[0] if g[1] == k else g[1]
            if not linia[inny // 2] or korzen(inny // 2) == r:
                break
            E = prosty(inny)
            if E is None:
                break
            sasiedzi.append((dict(dl=E[2]), E[1]))
        else:
            f = _faza(dict(pts=np.array([a, b]), dl=dl), sasiedzi[0], sasiedzi[1])
            if f:
                wynik["fazy"].append(f)

    # --- otwory i współśrodkowe okręgi
    okregi = [p for p in unikalne if p["typ"] == "CIRCLE"]
    wspol = set()
    for g in _klastry([p["srodek"] for p in okregi], STYK_TOL):
        fis = sorted({round(okregi[k]["fi"], 2) for k in g})
        if len(fis) >= 2:
            wynik["wspolsrodkowe"].append((fis, okregi[g[0]]["srodek"]))
            wspol.update(g)
    wynik["okregi"] = [(p["fi"], p["srodek"], p["reczny"], k in wspol) for k, p in enumerate(okregi)]

    # --- elementy dorysowane ręcznie (warstwa "0"), gdy reszta pochodzi z eksportu Inventora
    z_eksportu = [p for p in unikalne if not p["reczny"]]
    reczne = [p for p in unikalne if p["reczny"]]
    if z_eksportu and reczne:
        wynik["reczne"] = len(reczne)
        obrys = np.vstack([p["pts"] for p in z_eksportu])
        lo, hi = obrys.min(axis=0), obrys.max(axis=0)
        for p in reczne:
            wyst = float(max((lo - p["pts"].min(axis=0)).max(), (p["pts"].max(axis=0) - hi).max(), 0))
            if wyst > WYSTAJE_MAX:
                wynik["wystaje"].append((wyst, tuple(p["pts"][0])))
    return wynik


# ---------------------------------------------------------------- nakładka DXF na rysunek PDF (1:1)
# Rysunki z Inventora są WEKTOROWE: linie widoku są w PDF jako prawdziwe odcinki, nie piksele.
# 1. Szukanie widoku: (a) skupiska stykających się linii, których obrys ma wymiary DXF w skali widoku
#    (skale z opisów "( 1 : 10 )", a bez opisów — typowe skale), (b) gdy to zawiedzie (adnotacje
#    sklejone z widokiem, brak opisu, widok obrócony) — korelacja obrazu całej strony z obrazem DXF
#    (FFT) we wszystkich orientacjach i skalach.
# 2. Wstępne ustawienie każdej orientacji: korelacja FFT obrazu widoku z obrazem DXF (dowolne
#    przesunięcie), orientacje: obroty co 90°, lustro, plus kąt z minimalnego prostokąta otaczającego.
# 3. Dokładne dosunięcie: ICP sztywne (przesunięcie + mały obrót) na WEKTORACH — obraz służy tylko do
#    znalezienia przybliżonego położenia, pomiar odchyłek jest zawsze wektorowy.
# 4. Pomiar: odległość każdego punktu DXF od linii rysunku (DXF -> rysunek) oraz zamknięte elementy
#    wewnątrz widoku (otwory, wycięcia), których nie ma w DXF (rysunek -> DXF).
NAKLADKA_TOL_MM = 1.0   # mm — minimalna tolerancja zgodności DXF z rysunkiem
NAKLADKA_TOL_PT = 0.25  # pt na papierze — górna granica: dokładność zapisu widoku w PDF (zmierzona na PG4136348:
#                         ~0,12 pt), z zapasem; przy 1:k tolerancja <= max(1 mm, 0,25 pt * k) (1:10: 1 mm, 1:40: 3,5 mm)
NAKLADKA_TOL_PT_MIN = 0.06  # pt — dolna granica dla rysunków bardzo dokładnych (tolerancja dopasowana do szumu)
NAKLADKA_ZLY = 2.0      # mm — połowa punktów DXF dalej niż to (x tolerancja) = widok nie pasuje do DXF w ogóle
SKALE_RYSUNKU = (1, 1.5, 2, 2.5, 3, 4, 5, 6, 7, 7.5, 8, 10, 12, 12.5, 15, 20, 25, 30, 40, 50, 75, 100)
NAKLADKI_PDF = []       # (część, ścieżka PDF, strona, punkty DXF na stronie, odchyłki mm, tolerancja)


def tolerancja_nakladki(skala, r=None):
    """Tolerancja nakładki w mm dla widoku 1:skala. Z odchyłkami r (mm) dopasowana do dokładności tego
    rysunku: 4 x 90. percentyl odchyłek, w granicach [0,06 pt; 0,25 pt] * skala, nie mniej niż 1 mm. Dokładny
    rysunek -> ciasna tolerancja (brak fazy 5 mm przy 1:40 widać), rysunek "szumiący" -> górna granica."""
    pt_mm = skala * 25.4 / 72
    if r is None or not len(r):
        return max(NAKLADKA_TOL_MM, NAKLADKA_TOL_PT * pt_mm)
    szum = 4 * float(np.percentile(r, 90))
    return max(NAKLADKA_TOL_MM, min(NAKLADKA_TOL_PT * pt_mm, max(NAKLADKA_TOL_PT_MIN * pt_mm, szum)))


def _pdf_odcinki(strona):
    """Wszystkie odcinki rysunku (krzywe Beziera rozbite na 8 odcinków):
    [szer, x1, y1, x2, y2, wypełniony, kreskowany]. Wypełnione kształty to groty strzałek
    (przekroje, wymiary), kreskowane — linie ukryte, osie, linie gięcia."""
    out = []
    for d in strona.get_drawings():
        w = d.get("width") or 0
        wyp = 1.0 if d.get("fill") is not None else 0.0
        kresk = 1.0 if re.search(r"\d", (d.get("dashes") or "").split("]")[0]) else 0.0
        for it in d["items"]:
            if it[0] == "l":
                out.append((w, it[1].x, it[1].y, it[2].x, it[2].y, wyp, kresk))
            elif it[0] == "c":
                p = it[1:5]
                pts = [((1 - t) ** 3 * p[0].x + 3 * (1 - t) ** 2 * t * p[1].x + 3 * (1 - t) * t * t * p[2].x
                        + t ** 3 * p[3].x,
                        (1 - t) ** 3 * p[0].y + 3 * (1 - t) ** 2 * t * p[1].y + 3 * (1 - t) * t * t * p[2].y
                        + t ** 3 * p[3].y) for t in np.linspace(0, 1, 9)]
                out += [(w, *a, *b, wyp, kresk) for a, b in zip(pts, pts[1:])]
            elif it[0] == "re":
                r = it[1]
                c = [(r.x0, r.y0), (r.x1, r.y0), (r.x1, r.y1), (r.x0, r.y1)]
                out += [(w, *a, *b, wyp, kresk) for a, b in zip(c, c[1:] + c[:1])]
            elif it[0] == "qu":
                q = it[1]
                c = [(q.ul.x, q.ul.y), (q.ur.x, q.ur.y), (q.lr.x, q.lr.y), (q.ll.x, q.ll.y)]
                out += [(w, *a, *b, wyp, kresk) for a, b in zip(c, c[1:] + c[:1])]
    return np.array(out) if out else np.zeros((0, 7))


def _klasy_linii(seg):
    """Klasy grubości linii obecne na stronie (>= 4 odcinki), od najgrubszej. Linie widoczne
    konturu to zwykle jedna z grubszych klas (Inventor: 0,54 pt), wymiary cieńsze."""
    if not len(seg):
        return []
    szer, ile = np.unique(np.round(seg[:, 0], 2), return_counts=True)
    return [float(w) for w, n in sorted(zip(szer, ile), reverse=True) if n >= 4 and w > 0]


def _skupiska(seg, tol=0.8):
    """Grupy stykających się odcinków (do tol pt) — każdy widok rysunku to osobne skupisko."""
    n = len(seg)
    rodzic = list(range(n))

    def korzen(i):
        while rodzic[i] != i:
            rodzic[i] = rodzic[rodzic[i]]
            i = rodzic[i]
        return i

    konce = [tuple(p) for p in np.vstack([seg[:, 1:3], seg[:, 3:5]])]
    for g in _klastry(konce, tol):
        for k in g[1:]:
            a, b = korzen(g[0] % n), korzen(k % n)
            if a != b:
                rodzic[a] = b
    grupy = {}
    for i in range(n):
        grupy.setdefault(korzen(i), []).append(i)
    return [seg[v] for v in grupy.values() if len(v) >= 4]


def _zamkniete(g, tol=0.3):
    """Czy skupisko odcinków tworzy zamknięte pętle (otwór, wycięcie) — każdy węzeł ma parzysty
    stopień. Osie otworów, odnośniki, linie wymiarowe mają wolne końce."""
    konce = [tuple(p) for p in np.vstack([g[:, 1:3], g[:, 3:5]])]
    return all(len(k) % 2 == 0 for k in _klastry(konce, tol))


def _prostokat_min(pts):
    """Minimalny prostokąt otaczający (wypukła otoczka + obracanie): (boki rosnąco, kąt dłuższego boku)."""
    spts = sorted(set(map(tuple, np.round(pts, 3))))
    if len(spts) < 3:
        d = np.ptp(np.array(spts), axis=0) if spts else np.zeros(2)
        return np.sort(d), 0.0

    def cr(o, a, b):
        return (a[0] - o[0]) * (b[1] - o[1]) - (a[1] - o[1]) * (b[0] - o[0])

    lo, up = [], []
    for q in spts:
        while len(lo) >= 2 and cr(lo[-2], lo[-1], q) <= 0:
            lo.pop()
        lo.append(q)
    for q in reversed(spts):
        while len(up) >= 2 and cr(up[-2], up[-1], q) <= 0:
            up.pop()
        up.append(q)
    H = np.array(lo[:-1] + up[:-1])
    best = None
    for i in range(len(H)):
        d = H[(i + 1) % len(H)] - H[i]
        n = math.hypot(*d)
        if n < 1e-9:
            continue
        u = d / n
        a = H @ u
        b = H @ np.array([-u[1], u[0]])
        pole = np.ptp(a) * np.ptp(b)
        if best is None or pole < best[0]:
            kat = math.atan2(u[1], u[0]) if np.ptp(a) >= np.ptp(b) else math.atan2(u[0], -u[1])
            best = (pole, np.sort([np.ptp(a), np.ptp(b)]), kat)
    return best[1], best[2]


_SIATKA_ODCINKOW = [None, None]   # (seg, siatka) ostatniego zbioru — ICP pyta wiele razy o ten sam


def _odl_do_odcinkow(P, seg):
    """Odległość każdego punktu P do najbliższego odcinka seg ([., x1, y1, x2, y2, ...]) + ten punkt.
    Dla dużych zbiorów (blacha z setkami otworów) — siatka kubełkowa: najpierw odcinki z sąsiednich
    oczek; punkt, dla którego najbliższy znaleziony jest dalej niż oczko, liczony dokładnie po wszystkich."""
    if len(P) * len(seg) > 3_000_000 and len(seg) > 500:
        return _odl_siatka(P, seg)
    return _odl_brute(P, seg)


def _odl_siatka(P, seg):
    if _SIATKA_ODCINKOW[0] is not seg:
        A, B = seg[:, 1:3], seg[:, 3:5]
        lo = np.minimum(A, B).min(0)
        rozm = np.maximum(A, B).max(0) - lo
        C = max(2.0, float(max(rozm)) / 150)
        i0 = np.floor((np.minimum(A, B) - lo) / C).astype(int)
        i1 = np.floor((np.maximum(A, B) - lo) / C).astype(int)
        kubly = {}
        for k in range(len(seg)):
            for gx in range(i0[k, 0], i1[k, 0] + 1):
                for gy in range(i0[k, 1], i1[k, 1] + 1):
                    kubly.setdefault((gx, gy), []).append(k)
        _SIATKA_ODCINKOW[:] = [seg, (lo, C, {kl: np.array(v) for kl, v in kubly.items()}, {})]
    lo, C, kubly, sasiedztwo = _SIATKA_ODCINKOW[1]
    odl, naj = np.full(len(P), np.inf), np.zeros((len(P), 2))
    oczka = np.floor((P - lo) / C).astype(int)
    klucze, odwr = np.unique(oczka, axis=0, return_inverse=True)
    odwr = odwr.ravel()
    kolejnosc = np.argsort(odwr, kind="stable")
    granice = np.searchsorted(odwr[kolejnosc], np.arange(len(klucze) + 1))
    for u, (gx, gy) in enumerate(klucze):
        idx = kolejnosc[granice[u]:granice[u + 1]]
        kand = sasiedztwo.get((gx, gy))
        if kand is None:
            listy = [kubly[(gx + dx, gy + dy)] for dx in (-1, 0, 1) for dy in (-1, 0, 1) if (gx + dx, gy + dy) in kubly]
            kand = np.unique(np.concatenate(listy)) if listy else np.zeros(0, dtype=int)
            sasiedztwo[(gx, gy)] = kand
        if len(kand):
            o, q = _odl_brute(P[idx], seg[kand])
            odl[idx], naj[idx] = o, q
    daleko = ~(odl <= C)          # najbliższy spoza sąsiedztwa mógłby być bliżej — liczymy dokładnie
    if daleko.any():
        o, q = _odl_brute(P[daleko], seg)
        odl[daleko], naj[daleko] = o, q
    return odl, naj


def _odl_brute(P, seg):
    A, B = seg[:, 1:3], seg[:, 3:5]
    d = B - A
    dd = (d * d).sum(1)
    dd[dd == 0] = 1e-12
    odl, naj = np.empty(len(P)), np.empty((len(P), 2))
    blok = max(1, 2_000_000 // max(1, len(seg)))
    for i in range(0, len(P), blok):
        p = P[i:i + blok, None, :]
        t = np.clip(((p - A) * d).sum(2) / dd, 0, 1)
        q = A + d * t[..., None]
        r = np.hypot(q[..., 0] - p[..., 0], q[..., 1] - p[..., 1])
        j = r.argmin(1)
        odl[i:i + blok] = r[np.arange(len(j)), j]
        naj[i:i + blok] = q[np.arange(len(j)), j]
    return odl, naj


def _probki_dxf(path):
    """Punkty co ~1 mm wzdłuż elementów tnących DXF (te same warstwy co analiza konturu)."""
    pts = []
    for p in _prymitywy_dxf(ezdxf.readfile(path)):
        for a, b in zip(p["pts"], p["pts"][1:]):
            n = max(1, int(math.dist(a, b)))
            pts.append(a + (b - a) * np.linspace(0, 1, n, endpoint=False)[:, None])
        pts.append(p["pts"][-1:])
    return np.vstack(pts) if pts else np.zeros((0, 2))


def _orientacje(P, cel=None):
    """Macierze 2x2 orientacji DXF (bez skali): obroty co 90°, lustro, plus kąt wyrównujący
    minimalne prostokąty otaczające DXF i widoku (widok obrócony o dowolny kąt)."""
    katy = {0.0, 90.0, 180.0, 270.0}
    out = []
    if cel is not None and len(cel):
        _, a_dxf = _prostokat_min(P)
        _, a_wid = _prostokat_min(np.vstack([cel[:, 1:3], cel[:, 3:5]]) * np.array([1, -1]))
        for lustro in (False, True):
            a0 = math.pi - a_dxf if lustro else a_dxf
            baza = math.degrees(a_wid - a0) % 90
            if 0.5 < baza < 89.5:
                for k in range(4):
                    out.append((lustro, (baza + 90 * k) % 360))
    for lustro in (False, True):
        for k in katy:
            out.append((lustro, k))
    macierze = []
    for lustro, kat in out:
        c, s_ = math.cos(math.radians(kat)), math.sin(math.radians(kat))
        M = np.array([[c, -s_], [s_, c]]) @ (np.diag([-1.0, 1.0]) if lustro else np.eye(2))
        macierze.append((lustro, kat, M))
    return macierze


def _raster(pts, x0, y0, krok, ksztalt):
    """Punkty (pt) -> siatka 0/1 o oczku krok, początek (x0, y0)."""
    img = np.zeros(ksztalt, dtype=np.float32)
    ij = np.floor((pts - (x0, y0)) / krok).astype(int)
    ok = (ij[:, 0] >= 0) & (ij[:, 0] < ksztalt[1]) & (ij[:, 1] >= 0) & (ij[:, 1] < ksztalt[0])
    img[ij[ok, 1], ij[ok, 0]] = 1.0
    return img


def _probki_odcinkow(seg, krok):
    if not len(seg):
        return np.zeros((0, 2))
    A, B = seg[:, 1:3], seg[:, 3:5]
    n = np.maximum(2, (np.hypot(*(B - A).T) / krok).astype(int) + 2)
    return np.vstack([np.linspace(a, b, k) for a, b, k in zip(A, B, n)])


def _wstepne_ustawienia(P, cel, skala, okno, ile=3, caly_arkusz=False):
    """Korelacja FFT: obraz linii rysunku w oknie vs obraz DXF w każdej orientacji -> najlepsze
    (ocena = odsetek punktów DXF trafiających w linie, M, t). Siatka max ~350-600 oczek na bok."""
    s = 72 / 25.4 / skala
    x0, y0, x1, y1 = okno
    krok = max(0.5, max(x1 - x0, y1 - y0) / (600 if caly_arkusz else 350))
    H, W = int((y1 - y0) / krok) + 2, int((x1 - x0) / krok) + 2
    tlo = _raster(_probki_odcinkow(cel, krok / 2), x0, y0, krok, (H, W))
    pogr = tlo.copy()                                     # pogrubienie o 1 oczko: tolerancja położenia
    pogr[1:, :] = np.maximum(pogr[1:, :], tlo[:-1, :])
    pogr[:-1, :] = np.maximum(pogr[:-1, :], tlo[1:, :])
    tlo = pogr.copy()
    pogr[:, 1:] = np.maximum(pogr[:, 1:], tlo[:, :-1])
    pogr[:, :-1] = np.maximum(pogr[:, :-1], tlo[:, 1:])
    F = np.fft.rfft2(pogr)
    rzadkie = P[::max(1, len(P) // 3000)]
    wyniki = []
    for lustro, kat, M0 in _orientacje(P, None if caly_arkusz else cel):
        M = np.diag([s, -s]) @ M0                          # mm -> pt, oś y strony w dół
        X = rzadkie @ M.T
        mn = X.min(0)
        rozm = X.max(0) - mn
        if rozm[0] > (x1 - x0) + 2 * krok or rozm[1] > (y1 - y0) + 2 * krok:
            continue
        szablon = _raster(X - mn, 0, 0, krok, (H, W))
        ile_pkt = szablon.sum()
        if ile_pkt < 5:
            continue
        kor = np.fft.irfft2(F * np.conj(np.fft.rfft2(szablon)), s=(H, W))
        hs, ws = int(rozm[1] / krok) + 1, int(rozm[0] / krok) + 1
        kor = kor[:max(1, H - hs + 1), :max(1, W - ws + 1)]
        dy, dx = np.unravel_index(int(kor.argmax()), kor.shape)
        t = np.array([x0 + dx * krok, y0 + dy * krok]) - mn
        wyniki.append((float(kor[dy, dx] / ile_pkt), lustro, kat, M, t))
    wyniki.sort(key=lambda w: -w[0])
    # Orientacje o podobnej korelacji (np. prostokąt z jedną małą fazą: wszystkie 8 wyglądają w zgrubnej
    # siatce tak samo) sprawdzamy WSZYSTKIE — o wyborze decyduje dopiero dokładny pomiar wektorowy.
    if wyniki:
        wyniki = [w for w in wyniki if w[0] >= wyniki[0][0] - 0.1][:16] or wyniki[:ile]
    # doszukanie przesunięcia na drobnej siatce (~0,25 pt): zgrubne oczko bywa większe niż odstęp
    # linii podwójnych (krawędź ukosu 7 mm obok obrysu) — ICP startujące o oczko obok łapie złą linię
    out = []
    for ocena, lustro, kat, M, t in wyniki:
        out.append((ocena, lustro, kat, M, t + _doszukaj(rzadkie @ M.T + t, cel, krok)))
    return out


def _doszukaj(X, cel, krok):
    """Najlepsze przesunięcie X w zakresie ±krok na drobnej siatce (krok/5, min. 0,1 pt): odsetek
    punktów trafiających w (pogrubione o 1 oczko) linie rysunku. Zwraca poprawione przesunięcie."""
    f = max(0.1, krok / 5)
    lo, hi = X.min(0) - 2 * krok, X.max(0) + 2 * krok
    H, W = int((hi[1] - lo[1]) / f) + 2, int((hi[0] - lo[0]) / f) + 2
    if H * W > 6_000_000:
        f = math.sqrt((hi[1] - lo[1]) * (hi[0] - lo[0]) / 6_000_000)
        H, W = int((hi[1] - lo[1]) / f) + 2, int((hi[0] - lo[0]) / f) + 2
    siatka = _raster(_probki_odcinkow(_cel_w_oknie(cel, (lo[0], lo[1], hi[0], hi[1]), 0), f / 2), lo[0], lo[1], f, (H, W))
    pogr = siatka.copy()
    pogr[1:, :] = np.maximum(pogr[1:, :], siatka[:-1, :])
    pogr[:-1, :] = np.maximum(pogr[:-1, :], siatka[1:, :])
    siatka = pogr.copy()
    pogr[:, 1:] = np.maximum(pogr[:, 1:], siatka[:, :-1])
    pogr[:, :-1] = np.maximum(pogr[:, :-1], siatka[:, 1:])
    n = int(krok / f) + 1
    przes = np.arange(-n, n + 1) * f
    DX, DY = np.meshgrid(przes, przes)
    DX, DY = DX.ravel(), DY.ravel()
    ij = np.floor((X - lo) / f).astype(int)
    trafienia = np.zeros(len(DX))
    for k in range(len(DX)):
        jj = ij[:, 0] + int(round(DX[k] / f))
        ii = ij[:, 1] + int(round(DY[k] / f))
        ok = (ii >= 0) & (ii < H) & (jj >= 0) & (jj < W)
        trafienia[k] = pogr[ii[ok], jj[ok]].sum()
    najlepsze = np.flatnonzero(trafienia >= trafienia.max() - 1e-9)
    k = najlepsze[np.argmin(np.hypot(DX[najlepsze], DY[najlepsze]))]   # remis -> najmniejsze przesunięcie
    return np.array([DX[k], DY[k]])


def _icp(X0, cel, iteracje=40, obrot=True):
    """Sztywne ICP na wektorach: przesunięcie + (opcjonalnie) mały obrót, łącznie do 3°.
    X0 — punkty DXF już ustawione na stronie. Zwraca (Mc, tc) do X = X0 @ Mc.T + tc."""
    Mc, tc = np.eye(2), np.zeros(2)
    kat_suma = 0.0
    X = X0
    for it in range(iteracje):
        r, naj = _odl_do_odcinkow(X, cel)
        # próg odrzucania maleje od ~3 pt do 0,5 pt: punkty lekko odsunięte (start z dokładnością oczka
        # siatki) są najpierw PRZYCIĄGANE, a nie wyrzucane — inaczej wydłużony prostokąt zostawał
        # przesunięty wzdłuż długiego boku (krótkie boki odrzucone jako "obce")
        ok = r < max(3 * np.median(r), 3.0 * 0.75 ** it, 0.5)
        if ok.sum() < 3:
            break
        a, b = X[ok], naj[ok]
        ca, cb = a.mean(0), b.mean(0)
        fi = 0.0
        if obrot:
            u, v = a - ca, b - cb
            fi = math.atan2(float((u[:, 0] * v[:, 1] - u[:, 1] * v[:, 0]).sum()), float((u * v).sum()))
            if abs(kat_suma + fi) > math.radians(3):
                fi = 0.0
        kat_suma += fi
        Rm = np.array([[math.cos(fi), -math.sin(fi)], [math.sin(fi), math.cos(fi)]])
        t = cb - ca @ Rm.T
        X = X @ Rm.T + t
        Mc, tc = Rm @ Mc, Rm @ tc + t
        if math.hypot(*(cb - ca)) < 1e-4 and abs(fi) < 1e-7:
            break
    return Mc, tc


def _dopasuj(P, cel, skala, okno, caly_arkusz=False):
    """Najlepsze ułożenie DXF na liniach `cel` (pt) w `okno`. Zwraca słownik albo None."""
    s = 72 / 25.4 / skala
    geste = P[::max(1, len(P) // 600)]
    rzadkie = P[::max(1, len(P) // 200)]
    tol = tolerancja_nakladki(skala)

    def blisko(X):
        """Tylko linie rysunku w okolicy ułożonego DXF (na całej stronie bywa ich dziesiątki tysięcy)."""
        lo_, hi_ = X.min(0), X.max(0)
        z = 0.05 * float(max(hi_ - lo_)) + 10
        return _cel_w_oknie(cel, (lo_[0], lo_[1], hi_[0], hi_[1]), z)

    wstepne = []
    for ocena0, lustro, kat, M, t in _wstepne_ustawienia(P, cel, skala, okno, caly_arkusz=caly_arkusz):
        if ocena0 < 0.3:
            continue                                   # mniej niż 30% DXF trafia w linie — to nie ten widok
        lok = blisko(rzadkie @ M.T + t)
        if not len(lok):
            continue
        Mc, tc = _icp(rzadkie @ M.T + t, lok, iteracje=12)          # szybki etap dla każdej orientacji
        r, _ = _odl_do_odcinkow(rzadkie @ M.T @ Mc.T + Mc @ t + tc, lok)
        wstepne.append((int((r / s > tol).sum()), float(np.minimum(r / s, 20).mean()), lustro, Mc @ M,
                        Mc @ t + tc, ocena0))
    wstepne.sort(key=lambda k: (k[0], k[1]))
    kandydaci = []
    for _, _, lustro, M, t, ocena0 in wstepne[:3]:                   # pełny etap dla trzech najlepszych
        lok = blisko(geste @ M.T + t)
        Mc, tc = _icp(geste @ M.T + t, lok)
        r, _ = _odl_do_odcinkow(geste @ M.T @ Mc.T + Mc @ t + tc, lok)
        kandydaci.append((int((r / s > tol).sum()), float(np.minimum(r / s, 20).mean()), lustro, Mc @ M,
                          Mc @ t + tc, ocena0))
    # wariant bez lustra zawsze w pełnym etapie (reguła "lustro tylko, gdy bez lustra nie pasuje")
    if wstepne and all(k[2] for k in wstepne[:3]):
        bez = [k for k in wstepne if not k[2]]
        if bez:
            _, _, lustro, M, t, ocena0 = bez[0]
            lok = blisko(geste @ M.T + t)
            Mc, tc = _icp(geste @ M.T + t, lok)
            r, _ = _odl_do_odcinkow(geste @ M.T @ Mc.T + Mc @ t + tc, lok)
            kandydaci.append((int((r / s > tol).sum()), float(np.minimum(r / s, 20).mean()), lustro, Mc @ M,
                              Mc @ t + tc, ocena0))
    if not kandydaci:
        return None
    # najpierw najmniej punktów poza tolerancją (rozstrzyga orientację przy małych asymetriach), potem średnia
    kandydaci.sort(key=lambda k: (k[0], k[1]))
    best = kandydaci[0]
    bez_lustra = [k for k in kandydaci if not k[2]]
    if best[2] and bez_lustra and not best[0] < 0.5 * bez_lustra[0][0]:
        # Lustro tylko, gdy pasuje WYRAŹNIE lepiej (o połowę mniej punktów poza tolerancją). Część
        # symetryczna — także z wadą w jednym narożniku — pasuje bez lustra tak samo dobrze, a wtedy
        # podgląd i współrzędne różnic mają być po tej samej stronie co w DXF.
        best = bez_lustra[0]
    _, ocena, lustro, M, t, korelacja = best
    Q = P @ M.T + t
    r, _ = _odl_do_odcinkow(Q, blisko(Q))
    obrot = math.degrees(math.atan2(-M[1, 0], M[0, 0])) % 360      # kierunek osi X DXF na stronie (oś y w górę)
    return dict(ocena=float(np.minimum(r / s, 20).mean()), mediana=float(np.median(r)) / s, Q=Q, r=r / s,
                lustro=lustro, obrot=round(obrot, 1), skala=skala, korelacja=korelacja)


def _etykiety_skal(tekst):
    """Skale widoków z tekstu rysunku: "( 1 : 10 )", "SKALA 1:10", "Scale 1 : 10", "1:2,5"."""
    out = set()
    for m in re.finditer(r"(?:\(\s*|SKALA\s*:?\s*|SCALE\s*:?\s*|M[ÅA]LESTOKK\s*:?\s*)1\s*:\s*([\d]+(?:[,\.]\d+)?)",
                         tekst, re.I):
        try:
            k = float(m.group(1).replace(",", "."))
        except ValueError:
            continue
        if 0.1 <= k <= 500:
            out.add(k)
    return sorted(out)


def _cel_w_oknie(baza, okno, zapas):
    x0, y0, x1, y1 = okno
    m = ((baza[:, [1, 3]].max(1) >= x0 - zapas) & (baza[:, [1, 3]].min(1) <= x1 + zapas)
         & (baza[:, [2, 4]].max(1) >= y0 - zapas) & (baza[:, [2, 4]].min(1) <= y1 + zapas))
    return baza[m]


def nakladka(dxf_path, pdf_path):
    """Porównanie 1:1 konturu DXF z widokiem na rysunku PDF. Zwraca None, gdy na rysunku nie ma
    widoku pasującego do DXF (wtedy DXF jest NIEZWERYFIKOWANY), inaczej słownik z wynikiem."""
    P = _probki_dxf(dxf_path)
    if len(P) < 10:
        return None
    rozm, _ = _prostokat_min(P)
    if rozm[0] < 1:
        return None
    best = None
    doc = fitz.open(pdf_path)
    try:
        strony = []
        for nr, strona in enumerate(doc):
            seg = _pdf_odcinki(strona)
            baza = seg[(seg[:, 5] == 0) & (seg[:, 6] == 0)] if len(seg) else seg   # bez strzałek i kresek
            etykiety = _etykiety_skal(strona.get_text())
            strony.append((nr, strona.rect, baza, etykiety))

        def sprobuj(nr, k, cel, okno, zrodlo):
            nonlocal best
            if not len(cel):
                return
            wynik = _dopasuj(P, cel, k, okno, caly_arkusz=zrodlo == "strona")
            if wynik and (best is None or wynik["ocena"] < best["ocena"]):
                wynik.update(strona=nr, cel=cel, zrodlo=zrodlo, skala_opisana=k in etykiety_str[nr])
                best = wynik

        etykiety_str = {nr: et for nr, _, _, et in strony}
        # (a) skupiska linii o wymiarach DXF w skali widoku
        kandydaci = []
        for nr, rect, baza, etykiety in strony:
            skale = etykiety or SKALE_RYSUNKU
            for w in _klasy_linii(baza):
                klasa = baza[np.round(baza[:, 0], 2) == w]
                for g in _skupiska(klasa):
                    bb, _ = _prostokat_min(np.vstack([g[:, 1:3], g[:, 3:5]]))
                    for k in skale:
                        iloraz = bb * k * 25.4 / 72 / rozm
                        if np.all(iloraz > 0.97) and np.all(iloraz < 1.4):  # widok może mieć doklejone adnotacje
                            kandydaci.append((float(np.abs(iloraz - 1).sum()), nr, k, g, klasa))
        kandydaci.sort(key=lambda c: c[0])
        for _, nr, k, g, klasa in kandydaci[:4]:
            okno = (g[:, [1, 3]].min(), g[:, [2, 4]].min(), g[:, [1, 3]].max(), g[:, [2, 4]].max())
            zapas = 0.05 * max(okno[2] - okno[0], okno[3] - okno[1]) + 2
            okno = (okno[0] - zapas, okno[1] - zapas, okno[2] + zapas, okno[3] + zapas)
            sprobuj(nr, k, _cel_w_oknie(klasa, okno, 0), okno, "widok")
            if best and best["ocena"] < 0.3 * tolerancja_nakladki(k):
                break
        # (b) zapas: korelacja obrazu całej strony, każda klasa linii i skala
        if best is None or best["mediana"] > 0.5 * tolerancja_nakladki(best["skala"]):
            for nr, rect, baza, etykiety in strony:
                for k in (etykiety or SKALE_RYSUNKU):
                    rozm_pt = rozm * 72 / 25.4 / k
                    if rozm_pt[1] > max(rect.width, rect.height) or rozm_pt[0] < 15:
                        continue                          # DXF w tej skali nie mieści się albo jest za mały
                    for w in _klasy_linii(baza):
                        klasa = baza[np.round(baza[:, 0], 2) == w]
                        okno = (rect.x0, rect.y0, rect.x1, rect.y1)
                        sprobuj(nr, k, klasa, okno, "strona")
    finally:
        doc.close()
    if best is None or best["korelacja"] < 0.3:
        return None
    k = best["skala"]
    tol = tolerancja_nakladki(k, best["r"])
    Q, r, s = best["Q"], best["r"], 72 / 25.4 / k
    best["tol"] = tol
    # punkty DXF, których nie ma na rysunku
    best["rozne_dxf"] = _grupuj_punkty(P[r > tol], r[r > tol])
    # zamknięte elementy wewnątrz widoku (otwory, wycięcia), których nie ma w DXF
    best["brak_w_dxf"] = []
    lo, hi = Q.min(0), Q.max(0)
    cel = best.pop("cel")
    wnetrze = []
    for g in _skupiska(cel):
        glo = np.minimum(g[:, 1:3], g[:, 3:5]).min(0)
        ghi = np.maximum(g[:, 1:3], g[:, 3:5]).max(0)
        if np.all(glo > lo + 0.5) and np.all(ghi < hi - 0.5) and _zamkniete(g):
            wnetrze.append(g)
    if wnetrze:
        A, *_ = np.linalg.lstsq(np.c_[P, np.ones(len(P))], Q, rcond=None)    # DXF -> strona
        # raster DXF (oczko = tol/3, pogrubiony o 3 oczka ~ tol): "czy w odległości tol jest DXF?"
        c = tol * s / 3
        lo_r = Q.min(0) - 5 * c
        Hr, Wr = int(np.ptp(Q[:, 1]) / c) + 11, int(np.ptp(Q[:, 0]) / c) + 11
        Qg = np.vstack([np.linspace(a, b, max(2, int(math.dist(a, b) / (c / 2)) + 1)) for a, b in zip(Q, Q[1:])])
        jest = _raster(Qg, lo_r[0], lo_r[1], c, (Hr, Wr)) > 0
        for _ in range(3):
            p_ = jest.copy()
            p_[1:, :] |= jest[:-1, :]
            p_[:-1, :] |= jest[1:, :]
            p_[:, 1:] |= jest[:, :-1]
            p_[:, :-1] |= jest[:, 1:]
            jest = p_
        for g in wnetrze:
            probki = _probki_odcinkow(g, s)
            ij = np.floor((probki - lo_r) / c).astype(int)
            w_oknie = (ij[:, 0] >= 0) & (ij[:, 0] < Wr) & (ij[:, 1] >= 0) & (ij[:, 1] < Hr)
            trafione = np.zeros(len(probki), dtype=bool)
            trafione[w_oknie] = jest[ij[w_oknie, 1], ij[w_oknie, 0]]
            if trafione.mean() > 0.5:
                continue                                   # element jest w DXF
            srodek = np.linalg.solve(A[:2].T, probki.mean(0) - A[2])
            rozmiar = float(max(np.ptp(probki[:, 0]), np.ptp(probki[:, 1]))) / s
            best["brak_w_dxf"].append((rozmiar, tuple(srodek)))
        best["brak_w_dxf"].sort(reverse=True)
    return best


def _grupuj_punkty(punkty, odl):
    """Punkty odchyłek zgrupowane w miejsca (co 20 mm) -> [(max odchyłka, punkt)] malejąco."""
    if not len(punkty):
        return []
    miejsca = []
    for g in _klastry([tuple(p) for p in punkty], 20.0):
        k = max(g, key=lambda i: odl[i])
        miejsca.append((float(odl[k]), tuple(punkty[k])))
    return sorted(miejsca, reverse=True)


def ocen_nakladke(b, nak):
    """Wynik nakładki -> (błędy, uwagi, potwierdzony_1do1)."""
    bl, uw = [], []
    obrobka = "o" in b["flags"]
    tol = nak["tol"]
    widok = (f"strona {nak['strona'] + 1}, skala 1:{_mm(nak['skala'])}"
             + ("" if nak["skala_opisana"] else " (skala nieopisana na rysunku — dobrana)")
             + f", tolerancja {_mm(tol)} mm")
    if nak["mediana"] > NAKLADKA_ZLY * tol:
        if nak["korelacja"] >= 0.8:     # widok pewnie znaleziony (80% DXF trafia w linie), a kształt inny
            bl.append(f".DXF: nakładka — kontur DXF nie pokrywa się z widokiem rysunku ({widok}; połowa punktów "
                      f"dalej niż {nak['mediana']:.1f} mm od linii rysunku)")
        else:                           # najpewniej to nie ten widok — nie zgadujemy, oddajemy do ręcznego sprawdzenia
            uw.append(f".DXF: nakładka — na rysunku nie znaleziono widoku zgodnego z DXF (najbliższy: {widok}, połowa "
                      f"punktów dalej niż {nak['mediana']:.1f} mm) — porównanie 1:1 niemożliwe, sprawdź ręcznie")
        return bl, uw, False
    if nak["rozne_dxf"]:
        m = nak["rozne_dxf"]
        szara = m[0][0] <= 2 * tol          # tuż ponad tolerancję — może to być dokładność rysunku
        msg = (f".DXF: nakładka — DXF odbiega od rysunku 1:1 w {len(m)} miejscu(ach), maks. {m[0][0]:.1f} mm, "
               f"przy {_gdzie([p for _, p in m])} ({widok})")
        if obrobka:
            uw.append(msg + " — część z obróbką, możliwy naddatek")
        elif szara:
            uw.append(msg + f" — strefa szara (do {_mm(2 * tol)} mm): sprawdź na podglądzie NAKLADKI")
        else:
            bl.append(msg)
    if nak["brak_w_dxf"]:
        brak = nak["brak_w_dxf"]
        (uw if obrobka else bl).append(
            f".DXF: nakładka — na rysunku jest {len(brak)} zamknięty element (otwór/wycięcie, największy "
            f"~{brak[0][0]:.0f} mm), którego nie ma w DXF, przy {_gdzie([x[1] for x in brak])} ({widok})"
            + (" — część z obróbką: otwór wykonywany/obrabiany później (w DXF mniejszy albo brak)" if obrobka else ""))
    if nak["lustro"]:
        uw.append(f".DXF: nakładka — DXF pasuje do rysunku dopiero po ODBICIU LUSTRZANYM ({widok}); "
                  f"sprawdź stronę gięcia/ukosu")
    return bl, uw, not bl and not nak["rozne_dxf"]


def zapisz_nakladki_pdf(sciezka):
    """Zbiorczy PDF: strona rysunku z naniesionym konturem DXF (zielony = zgodny, czerwony = różnica)."""
    out = fitz.open()
    for etykieta, pdf_path, nr, Q, r, tol in NAKLADKI_PDF:
        with fitz.open(pdf_path) as src:
            out.insert_pdf(src, from_page=nr, to_page=nr)
        strona = out[-1]
        ksztalt = strona.new_shape()      # jeden kształt na stronę — tysiące kropek osobno trwałyby minuty
        for kolor, zly in (((0, 0.6, 0), False), ((0.9, 0, 0), True)):
            for i in range(len(Q) - 1):
                if (r[i] > tol) == zly and math.dist(Q[i], Q[i + 1]) < 3:
                    ksztalt.draw_line(Q[i], Q[i + 1])
            ksztalt.finish(color=kolor, width=1.6 if zly else 0.8)
        ksztalt.commit()
        strona.insert_text((20, 14), f"NAKLADKA DXF: {etykieta}  (zielony = zgodny z rysunkiem, czerwony = roznica "
                           f"> {tol:.1f} mm)", fontsize=9, color=(0.9, 0, 0))
    if len(out):
        out.save(sciezka)
    out.close()
    return sciezka if NAKLADKI_PDF else None


# skok gwintu zwykłego ISO -> otwór pod gwint = M - skok (M12 -> 10,2)
SKOK_GWINTU = {3: 0.5, 4: 0.7, 5: 0.8, 6: 1.0, 8: 1.25, 10: 1.5, 12: 1.75, 14: 2.0, 16: 2.0, 18: 2.5,
               20: 2.5, 22: 2.5, 24: 3.0, 27: 3.0, 30: 3.5, 36: 4.0, 42: 4.5, 48: 5.0}


def wiertlo_pod_gwint(d):
    return d - SKOK_GWINTU.get(round(d), 0.125 * d)


def _mm(x):
    return f"{x:.2f}".rstrip("0").rstrip(".").replace(".", ",")


def _gdzie(punkty, ile=3):
    s = ", ".join(f"({x:.1f}; {y:.1f})" for x, y in punkty[:ile])
    return s + (" …" if len(punkty) > ile else "")


def _lista_faz(fazy):
    licz = {}
    for n1, n2, _ in fazy:
        licz[(n1, n2)] = licz.get((n1, n2), 0) + 1
    return ", ".join(f"{_mm(a)}x{_mm(b)}" + (f" ({n} szt.)" if n > 1 else "") for (a, b), n in licz.items())


def ocen_kontur(b, kat, ana, tresc, nakladka_stan=None, nak=None):
    """Zamienia wynik dxf_analiza (+ wymiary z rysunku PDF, jeśli jest) na listy
    (błędy, uwagi). Każdy komunikat zaczyna się od ".DXF:", żeby trafił do kolumny Plik."""
    bl, uw = [], []
    ukos = "u" in b["flags"]
    obrobka = "o" in b["flags"]
    gieta = "p" in b["flags"]
    przymiar = (kat == "przygotowka" or b["part"].upper().endswith("PRZ")
                or re.search(r"PRZYMIAR|PRZYG", f"{b['title']} {b['desc']}", re.I))

    if ana.get("pusty"):
        bl.append(".DXF: kontur — w DXF nie ma żadnej geometrii na warstwach cięcia (pusty plik albo wszystko "
                  "na warstwach pomijanych: wymiary, osie, gięcie)")
    if ana.get("jednostki") in JEDNOSTKI_DXF:
        # UWAGA, nie BŁĄD: liczby w pliku i tak porównujemy z BOM w mm (gabaryt) — prawdziwy błąd skali
        # wyjdzie tam; to pole decyduje tylko, czy program CAM przy imporcie nie przeskaluje części.
        uw.append(f".DXF: kontur — nagłówek DXF mówi, że jednostką są {JEDNOSTKI_DXF[ana['jednostki']]} "
                  f"($INSUNITS={ana['jednostki']}), a nie milimetry — sprawdź, czy program wypalarki nie "
                  f"przeskaluje części przy imporcie")

    if ana["duble"] >= 3:
        bl.append(f".DXF: podwójne linie — {ana['duble']} elementów narysowanych 2x w tym samym miejscu "
                  f"(laser tnie dwa razy po tym samym śladzie)")
    elif ana["duble"]:
        uw.append(f".DXF: podwójne linie — {ana['duble']} zdublowany element (bywa z eksportu)")

    if ana["przerwy"]:
        d_max = max(d for d, _ in ana["przerwy"])
        bl.append(f".DXF: kontur — {len(ana['przerwy'])} przerw(a) w konturze (największa {_mm(d_max)} mm), "
                  f"przy {_gdzie([p for _, p in ana['przerwy']])}")
    if ana["wolne"]:
        (uw if przymiar else bl).append(
            f".DXF: kontur — {len(ana['wolne'])} wolnych końców linii (kontur niedomknięty albo linia "
            f"wystaje za narożnik), przy {_gdzie(ana['wolne'])}"
            + (" — przygotówka/przymiar: to mogą być linie trasowania" if przymiar else ""))
    if ana["rozgal"]:
        # BŁĄD także przy ukosowaniu i gięciu: wzorcowy DXF części "p u" (PG4136348) ma sam obrys — linie
        # ukosu/gięcia na warstwie cięcia laser by przeciął. Wyjątek: przygotówki (linie trasowania).
        (uw if przymiar else bl).append(f".DXF: kontur — {len(ana['rozgal'])} rozgałęzień (linia kończy się na innej linii albo 3+ "
                  f"linie w jednym punkcie), przy {_gdzie(ana['rozgal'])} — nieprzycięty narożnik po dorysowaniu "
                  f"fazy albo linia fazy/ukosu/gięcia na warstwie cięcia, którą laser wytnie"
                  + (" (część z ukosowaniem/gięciem: te linie mają być na osobnej warstwie albo usunięte)"
                     if ukos or gieta else ""))
    for fis, srodek in ana["wspolsrodkowe"]:
        msg = (f".DXF: otwory — współśrodkowe okręgi Ø{' / Ø'.join(_mm(f) for f in fis)} przy "
               f"{_gdzie([srodek])} — faza/pogłębienie otworu wyeksportowane do DXF, laser wytnie większy okrąg")
        (uw if obrobka else bl).append(msg)
    if ana["wystaje"] and not przymiar:
        w = max(ana["wystaje"])
        bl.append(f".DXF: dorysowane — element z warstwy 0 wystaje {w[0]:.0f} mm poza obrys części przy "
                  f"{_gdzie([w[1]])} (zbłąkana geometria, cięcie w powietrzu)")
    if ana["reczne"]:
        uw.append(f".DXF: dorysowane — {ana['reczne']} element(y) dorysowane ręcznie na warstwie 0 (do wglądu)")

    if not tresc:
        return bl, uw

    # --- otwory z DXF vs Ø na rysunku
    fi_rys = tresc.get("fi", []) + [2 * r for r in tresc.get("promienie", [])]

    def zwymiarowany(fi):
        return (any(abs(fi - f) <= OTWOR_TOL for f in fi_rys)
                or any(abs(fi - wiertlo_pod_gwint(d)) <= 0.3 or abs(fi - d) <= OTWOR_TOL
                       for d in tresc.get("gwinty", [])))

    if not obrobka and not przymiar:
        reczne_bez = sorted({round(fi, 2) for fi, _, reczny, wsp in ana["okregi"]
                             if reczny and not wsp and not zwymiarowany(fi)})
        eksp_bez = sorted({round(fi, 2) for fi, _, reczny, wsp in ana["okregi"]
                           if not reczny and not wsp and not zwymiarowany(fi)})
        # Bez łagodzenia przy potwierdzonej nakładce: tolerancja nakładki (>= 1 mm) jest grubsza niż
        # porównanie średnicy z wymiarem Ø (0,15 mm) — Ø15 zamiast Ø13 to tylko 1 mm na promieniu.
        potw = nakladka_stan == "potwierdzony"
        if reczne_bez:
            bl.append(f".DXF: otwory — dorysowany otwór Ø{', Ø'.join(_mm(f) for f in reczne_bez)} (warstwa 0) "
                      f"bez wymiaru Ø na rysunku PDF — produkcja nie wie o otworze"
                      + (" (nakładka: otwór jest narysowany na rysunku w tym miejscu)" if potw else ""))
        if eksp_bez and tresc.get("fi"):   # stare rysunki bez wymiarów Ø — nie sprawdzamy
            bl.append(f".DXF: otwory — otwór Ø{', Ø'.join(_mm(f) for f in eksp_bez)} z DXF nie występuje "
                      f"w wymiarach rysunku (Ø na rysunku: {', '.join(_mm(f) for f in sorted(set(tresc['fi'])))})"
                      + (" — różnica mniejsza niż tolerancja nakładki, ale średnica nie zgadza się z wymiarem"
                         if potw else ""))

    # --- fazy narożników z DXF vs notki faz na rysunku. Gdy nakładka 1:1 potwierdziła kontur, każdy
    # narożnik jest już sprawdzony geometrycznie, a notki "7 X 45° Chamfer" bywają fazami KRAWĘDZI
    # (ukos pod spaw w przekroju), których w konturze nie ma — wtedy tego porównania nie robimy.
    fazy_dxf = ana["fazy"]
    if nakladka_stan is not None:
        # Widok znaleziony: brak całej fazy albo faza w złym miejscu wychodzi w nakładce (kilka mm).
        # Notki bez żadnej podobnej fazy w konturze to zwykle fazy KRAWĘDZI (ukos pod spaw w przekroju) —
        # pomijamy. Ale faza w DXF PRAWIE jak w notce (np. 11 zamiast 10) jest poniżej tolerancji
        # nakładki — tę różnicę łapiemy tutaj, z tolerancją FAZA_TOL.
        # tylko notki stojące przy dopasowanym widoku rozwinięcia — notki przy przekrojach (A-A, C-C)
        # opisują fazy krawędzi pod spaw, nie narożniki konturu (wzorzec: PG4136348, "7,00 X 45° Chamfer")
        notki = []
        if nak is not None:
            lo, hi = nak["Q"].min(0), nak["Q"].max(0)
            zapas = np.array([60.0, 60.0])     # ~2 cm na papierze: notka z odnośnikiem stoi przy widoku
            notki = [f for f in tresc.get("notki_faz", []) if f["strona"] == nak["strona"]
                     and lo[0] - zapas[0] <= f["x"] <= hi[0] + zapas[0] and lo[1] - zapas[1] <= f["y"] <= hi[1] + zapas[1]]
        for f in notki:
            if max(f["nogi"]) <= FAZA_MIN_RYS:
                continue                       # łamanie krawędzi
            zgodne = [d for d in fazy_dxf if abs(d[0] - f["nogi"][0]) <= FAZA_TOL and abs(d[1] - f["nogi"][1]) <= FAZA_TOL]
            if zgodne:
                if f["ile"] and len(zgodne) < f["ile"]:
                    (uw if ukos else bl).append(
                        f".DXF: fazy — przy widoku rozwinięcia notka {f['tekst']} ({f['ile']} szt.), a w DXF takich "
                        f"faz narożnika jest {len(zgodne)}")
                continue
            blisko = [d for d in fazy_dxf if abs(d[0] - f["nogi"][0]) <= 3 and abs(d[1] - f["nogi"][1]) <= 3]
            if blisko:
                bl.append(f".DXF: fazy — na rysunku faza {f['tekst']}, a w DXF {_lista_faz(blisko)} przy "
                          f"{_gdzie([d[2] for d in blisko])} — różnica może być poniżej tolerancji nakładki, ale "
                          f"wymiar fazy się nie zgadza")
            else:
                (uw if ukos else bl).append(
                    f".DXF: fazy — przy widoku rozwinięcia jest notka fazy {f['tekst']}, a w konturze DXF nie ma "
                    f"takiej fazy narożnika" + (" (część z ukosowaniem — może to faza krawędzi)" if ukos else ""))
        return bl, uw
    lagodnie = ukos or obrobka

    def pasuje(n, m):
        return abs(n[0] - m[0]) <= FAZA_TOL and abs(n[1] - m[1]) <= FAZA_TOL

    rys = {}
    for f in tresc.get("fazy", []):
        if max(f["nogi"]) <= FAZA_MIN_RYS:
            continue  # łamanie krawędzi (np. 1x45° na otworze) — nie ma go w konturze
        r = rys.setdefault(f["nogi"], dict(teksty=[], ile=0, z_iloscia=False))
        r["teksty"].append(f["tekst"])
        r["ile"] += f["ile"] or 1
        r["z_iloscia"] |= f["ile"] is not None
    # Bez widoku nie wiadomo, czy notka dotyczy narożnika konturu, czy KRAWĘDZI (ukos, lemiesz) — tylko
    # UWAGA; i tak DXF bez nakładki nie dostanie statusu "ZGODNY 1:1".
    lagodnie = True
    dopisek = " (bez dopasowanego widoku nie wiadomo, czy to faza narożnika, czy krawędzi — sprawdź)"
    if rys:
        for nogi, r in rys.items():
            n = sum(1 for f in fazy_dxf if pasuje(f, nogi))
            if n == 0:
                (uw if lagodnie else bl).append(
                    f".DXF: fazy — na rysunku faza {r['teksty'][0]}, w DXF brak takiej fazy ("
                    + (f"fazy w DXF: {_lista_faz(fazy_dxf)}" if fazy_dxf else "w DXF nie ma żadnej fazy narożnika")
                    + ")" + dopisek)
            elif r["z_iloscia"] and n < r["ile"]:
                (uw if lagodnie else bl).append(
                    f".DXF: fazy — na rysunku {r['ile']}x faza {_mm(nogi[0])}x{_mm(nogi[1])}, w DXF tylko {n}" + dopisek)
        obce = [f for f in fazy_dxf if not any(pasuje(f, nogi) for nogi in rys)]
        if obce:
            uw.append(f".DXF: fazy — w DXF faza {_lista_faz(obce)} przy {_gdzie([f[2] for f in obce])}, "
                      f"której nie ma na rysunku (rysunek: {', '.join(r['teksty'][0] for r in rys.values())})")
    elif fazy_dxf and tresc.get("liczby"):
        nie = [f for f in fazy_dxf
               if not any(abs(f[0] - x) <= FAZA_TOL or abs(f[1] - x) <= FAZA_TOL for x in tresc["liczby"])]
        if nie:
            uw.append(f".DXF: fazy — faza {_lista_faz(nie)} z DXF nie jest zwymiarowana na rysunku PDF, "
                      f"przy {_gdzie([f[2] for f in nie])}")
    return bl, uw


# ---------------------------------------------------------------- geometria STP (3D)
_INV_APP = None
_INV_STARTED_BY_US = False
_STP_OSTRZEZONO = False  # ostrzeżenie o braku Inventora pokazujemy raz na całe uruchomienie


def get_inventor():
    """Zwraca instancję Inventor.Application przez COM — używa już uruchomionej (jeśli jest),
    inaczej odpala nową w tle (niewidoczną). Nie zamyka cudzej, już otwartej instancji na końcu."""
    global _INV_APP, _INV_STARTED_BY_US
    if _INV_APP is not None:
        return _INV_APP
    try:
        _INV_APP = win32com.client.GetActiveObject("Inventor.Application")  # dolacz do uruchomionej
    except Exception:
        _INV_APP = win32com.client.Dispatch("Inventor.Application")  # odpal nowa (w tle)
        try:
            _INV_APP.Visible = False
        except Exception:
            pass
        _INV_STARTED_BY_US = True
    return _INV_APP


@atexit.register
def _zamknij_inventor():
    if _INV_STARTED_BY_US and _INV_APP is not None:
        try:
            _INV_APP.Quit()
        except Exception:
            pass


def znajdz_stp(folder):
    """Szuka pliku STP dla projektu: najpierw '<nazwa_folderu_bez_dopiskow>.stp/.step'
    (wzorzec z archiwum: GE6725\\GE6725.stp), inaczej dowolny *.stp/*.step w folderze."""
    baza = re.sub(r"[\s\(\)].*$", "", os.path.basename(folder.rstrip("\\/")))
    for ext in (".stp", ".step", ".STP", ".STEP"):
        kand = os.path.join(folder, baza + ext)
        if os.path.isfile(kand):
            return kand
    try:
        for f in os.listdir(folder):
            if f.lower().endswith((".stp", ".step")):
                return os.path.join(folder, f)
    except OSError:
        pass
    return None


def stp_dims_mapa(stp_path):
    """Otwiera plik STP przez COM Inventora, zwraca (mapa, blad) gdzie mapa to
    {NAZWA_CZESCI_UPPER: (dx, dy, dz) w mm} z LOKALNEGO (nieobróconego) RangeBox definicji
    każdej liściastej części — odpornego na obrót/pozycję części w zespole.
    Jednostki bazowe Inventora to zawsze centymetry, stąd mnożenie *10.0 -> mm (nie wołamy
    UnitsOfMeasure.ConvertUnits przez surowe COM, żeby nie zgadywać numeru wyliczenia)."""
    global _STP_OSTRZEZONO
    if not STP_MODULE_OK:
        return {}, "brak modułu pywin32 (pip install pywin32) — kontrola STP wyłączona"
    try:
        oApp = get_inventor()
    except Exception as e:
        if not _STP_OSTRZEZONO:
            _STP_OSTRZEZONO = True
        return {}, f"nie udało się uruchomić/połączyć z Inventorem ({type(e).__name__}: {e})"

    mapa = {}
    oDoc = None
    try:
        oDoc = oApp.Documents.Open(stp_path, False)
    except TypeError:
        try:
            oDoc = oApp.Documents.Open(stp_path)
        except Exception as e:
            return {}, f"nie udało się otworzyć pliku STP ({type(e).__name__}: {e})"
    except Exception as e:
        return {}, f"nie udało się otworzyć pliku STP ({type(e).__name__}: {e})"

    try:
        try:
            # oDoc.ComponentDefinition rzuca AttributeError na "gołym" obiekcie Document
            # zwróconym przez Documents.Open (early-bound gen_py widzi tylko bazowy interfejs
            # Document, nie AssemblyDocument) - CastTo jest oficjalnym lekarstwem pywin32 na
            # dokładnie ten przypadek. Potwierdzone empirycznie 2026-07-18 na GE75206-S1.stp
            # (64 wystąpienia, iProperty "Part Number" poprawnie = numer części z BOM).
            oAsmDoc = win32com.client.CastTo(oDoc, "AssemblyDocument")
            oCompDef = oAsmDoc.ComponentDefinition
            oLeafOcc = oCompDef.Occurrences.AllLeafOccurrences  # wymusza blad tu, nie w petli, jesli to nie zespol
        except Exception:
            return {}, ("plik STP zaimportowany jako pojedyncza część (nie zespół) — Inventor "
                        "nie rozbił go na komponenty; sprawdzenie geometrii per-część pominięte "
                        "(sam plik nie jest błędem)")
        for oOcc in oLeafOcc:
            nazwa = None
            try:
                nazwa = oOcc.Definition.Document.PropertySets.Item(
                    "Design Tracking Properties").Item("Part Number").Value
            except Exception:
                pass
            if not nazwa:
                try:
                    nazwa = oOcc.Definition.Document.DisplayName
                except Exception:
                    pass
            if not nazwa:
                try:
                    nazwa = oOcc.Name.split(":")[0]
                except Exception:
                    continue
            try:
                box = oOcc.Definition.RangeBox
                dx = abs(box.MaxPoint.X - box.MinPoint.X) * 10.0
                dy = abs(box.MaxPoint.Y - box.MinPoint.Y) * 10.0
                dz = abs(box.MaxPoint.Z - box.MinPoint.Z) * 10.0
            except Exception:
                continue
            mapa[str(nazwa).strip().upper()] = (dx, dy, dz)
    except Exception as e:
        return mapa, f"błąd podczas czytania zespołu STP ({type(e).__name__}: {e})"
    finally:
        try:
            if oDoc:
                oDoc.Close(True)
        except Exception:
            pass
    return mapa, None


def _stp_klucz(nazwa):
    """Nazwa części z STP bez dopisków Inventora: rozszerzenie (.ipt/.stp), numer wystąpienia (":1")."""
    n = str(nazwa).strip().upper()
    n = re.sub(r"\.(IPT|IAM|STP|STEP)$", "", n)
    return re.sub(r"\s*:\s*\d+$", "", n).strip()


def stp_szukaj(stp_mapa, part):
    """Wymiary części z mapy STP: dokładna nazwa, potem bez dopisków (":1", ".ipt"), potem nazwa
    zawierająca numer części jako osobny człon (np. "PG4136348 - TOP BEND PLATE")."""
    part = part.strip().upper()
    if part in stp_mapa:
        return stp_mapa[part]
    for k, v in stp_mapa.items():
        if _stp_klucz(k) == part:
            return v
    wzor = re.compile(r"(?<![A-Z0-9])" + re.escape(part) + r"(?![A-Z0-9])")
    trafione = [v for k, v in stp_mapa.items() if wzor.search(k)]
    return trafione[0] if len(trafione) == 1 else None


# ---------------------------------------------------------------- pomocnicze do raportu xlsx
def etykieta_split(etykieta):
    """'[7.3.2] PG4104413 (WZM BOKU)' -> ('7.3.2', 'PG4104413', 'WZM BOKU')."""
    m = re.match(r"^\[(.*?)\]\s+(\S+)\s+\((.*)\)\s*$", etykieta)
    if m:
        return m.group(1), m.group(2), m.group(3)
    return "", etykieta, ""


def plik_i_reszta(opis):
    """'.STP: grubosc 12 vs BOM (...)' -> ('STP', 'grubosc 12 vs BOM (...)')."""
    m = re.match(r"^\.(DXF|PDF|STP)\s*:\s*(.*)$", opis, re.I)
    if m:
        return m.group(1).upper(), m.group(2)
    return "", opis


RODZAJ_WZORCE = [
    (re.compile(r"^nakładka —|nakładka —"), "DXF vs rysunek 1:1 (nakładka)"),
    (re.compile(r"^kontur —"), "kontur DXF (przerwy/rozgałęzienia)"),
    (re.compile(r"^fazy —"), "fazy DXF vs rysunek"),
    (re.compile(r"^otwory —"), "otwory DXF vs rysunek"),
    (re.compile(r"^podwójne linie —"), "podwójne linie DXF"),
    (re.compile(r"^dorysowane —"), "dorysowane ręcznie (warstwa 0)"),
    (re.compile(r"^BRAK (DXF|PDF)"), "plik brakujący"),
    (re.compile(r"geometria DXF|DXF zawiera więcej"), "geometria DXF"),
    (re.compile(r"bryła", re.I), "geometria STP (3D)"),  # prefiks ".STP:" juz wyciety do kolumny Plik
    (re.compile(r"grubo(ś|s)ć"), "grubość vs BOM"),
    (re.compile(r"materia(ł|l)"), "materiał vs BOM"),
    (re.compile(r"ilo(ś|s)ć|stempel"), "ilość/stempel"),
    (re.compile(r"flagi"), "flagi MOPUC"),
    (re.compile(r"tabliczka"), "tabliczka PDF"),
    (re.compile(r"prefiks"), "prefiks nazwy pliku"),
    (re.compile(r"bez pozycji w BOM"), "plik bez pozycji w BOM"),
]


def rodzaj_z_opisu(opis):
    for pat, nazwa in RODZAJ_WZORCE:
        if pat.search(opis):
            return nazwa
    return "inne"


def pos_key(pos):
    """Sortowanie POS naturalnie/numerycznie: '7.3.2' < '10' (nie leksykograficznie)."""
    klucz = []
    for p in re.split(r"[.\-]", pos or ""):
        try:
            klucz.append((0, float(p.replace(",", "."))))
        except ValueError:
            klucz.append((1, p))
    return klucz


def dodaj_wiersze(folder, lista, poziom):
    for etykieta_pelna in lista:
        if ": " in etykieta_pelna and etykieta_pelna.count("[") >= 1:
            etykieta, opis = etykieta_pelna.split(": ", 1)
        else:
            etykieta, opis = "", etykieta_pelna
        pos, part, tytul = etykieta_split(etykieta) if etykieta else ("", etykieta_pelna.split(":")[0], "")
        plik, reszta = plik_i_reszta(opis)
        rodzaj = "model 3D .stp (informacyjnie)" if poziom == "INFO STP" else rodzaj_z_opisu(reszta)
        WSZYSTKIE_WIERSZE.append(dict(
            folder=os.path.basename(folder.rstrip("\\/")),
            pos=pos, part=part, tytul=tytul, poziom=poziom,
            plik=plik, rodzaj=rodzaj, opis=reszta,
        ))


# ---------------------------------------------------------------- kontrola
def folder_niesprawdzony(folder, powod):
    """Folder, którego nie dało się sprawdzić, MUSI być w raporcie jako błąd — nigdy cicho pominięty."""
    nazwa = os.path.basename(folder.rstrip("\\/"))
    WSZYSTKIE_WIERSZE.append(dict(folder=nazwa, pos="", part="(cały folder)", tytul="", poziom="BŁĄD", plik="",
                                  rodzaj="folder niesprawdzony", opis=f"folder NIE został sprawdzony: {powod}"))
    PODSUMOWANIE_FOLDEROW.append(dict(folder=nazwa, sciezka=folder, ok=0, bledy=1, uwagi=0, pozycji=0, stp="-",
                                      stp_info=0, dxf_sprawdzone=0, dxf_bledne=[], dxf_uwagi=0, dxf_1do1=0,
                                      dxf_wszystkie=0, dxf_niezweryf=[], dxf_rozne=[], niesprawdzony=powod))
    return 1


def autotest():
    """Sprawdzenie na wzorcowej części, czy kontrola geometrii działa NA TYM KOMPUTERZE (wersje bibliotek,
    zbudowany .exe): dobry DXF musi przejść nakładkę 1:1, zepsuty (bez fazy, przesunięty otwór) — nie,
    przerwa w konturze musi zostać znaleziona. Wynik: dict(ok, opis)."""
    import tempfile
    import time
    t0 = time.time()
    try:
        with tempfile.TemporaryDirectory() as tmp:
            def dxf(nazwa, faza, otwor_x, przerwa=0.0):
                d = ezdxf.new("R2010", units=ezdxf.units.MM)
                m = d.modelspace()
                pkt = [(0, 0), (200 - faza, 0)] + ([(200, faza)] if faza else []) + [(200, 100), (0, 100)]
                for i, (a, b) in enumerate(zip(pkt, pkt[1:] + pkt[:1])):
                    if i == 0 and przerwa:
                        b = (b[0] - przerwa, b[1])
                    m.add_line(a, b, dxfattribs={"layer": "Visible (ISO)"})
                m.add_circle((otwor_x, 50), 10, dxfattribs={"layer": "Visible (ISO)"})
                sciezka = os.path.join(tmp, nazwa)
                d.saveas(sciezka)
                return sciezka

            dobry, zly, z_przerwa = dxf("dobry.dxf", 10, 50), dxf("zly.dxf", 0, 55), dxf("przerwa.dxf", 10, 50, 0.5)
            doc = fitz.open()
            strona = doc.new_page(width=842, height=595)
            sk = 72 / 25.4 / 2                                       # widok 1:2
            P = lambda x, y: fitz.Point(150 + x * sk, 400 - y * sk)  # noqa: E731
            pkt = [(0, 0), (190, 0), (200, 10), (200, 100), (0, 100)]
            for a, b in zip(pkt, pkt[1:] + pkt[:1]):
                strona.draw_line(P(*a), P(*b), width=0.54)
            strona.draw_circle(P(50, 50), 10 * sk, width=0.54)
            strona.draw_line(P(0, -15), P(200, -15), width=0.36)
            strona.insert_text((150, 120), "VIEW1 ( 1 : 2 )   4x n20   10 x 45°", fontsize=9)
            rys = os.path.join(tmp, "rysunek.pdf")
            doc.save(rys)
            doc.close()

            n = nakladka(dobry, rys)
            if not n or n["rozne_dxf"] or n["brak_w_dxf"] or n["r"].max() > 0.2 or n["lustro"]:
                return dict(ok=False, opis="poprawny DXF wzorcowy nie przeszedł nakładki 1:1")
            n = nakladka(zly, rys)
            if not n or not n["rozne_dxf"] or n["rozne_dxf"][0][0] < 4:
                return dict(ok=False, opis="zepsuty DXF wzorcowy (brak fazy, przesunięty otwór) nie został wykryty")
            if len(dxf_analiza(z_przerwa)["przerwy"]) != 1:
                return dict(ok=False, opis="przerwa 0,5 mm w konturze wzorcowym nie została wykryta")
            w = wymiary_rysunku("4x n20   10 x 45°")
            if 20.0 not in w["fi"] or not any(f["nogi"] == (10.0, 10.0) for f in w["fazy"]):
                return dict(ok=False, opis="odczyt wymiarów z tekstu rysunku nie działa")
    except Exception as e:
        return dict(ok=False, opis=f"wyjątek {type(e).__name__}: {e}")
    return dict(ok=True, opis=f"wzorcowe DXF sprawdzone poprawnie ({time.time() - t0:.1f} s)")


def sprawdz(folder):
    folder = folder.rstrip("\\/")
    print("\n" + "=" * 92)
    print(f"KONTROLA (GEO+STP): {folder}")
    if not GEO:
        print("  UWAGA: brak ezdxf (pip install ezdxf) — kontrola geometrii DXF wyłączona")
    if not STP_MODULE_OK:
        print("  UWAGA: brak pywin32 (pip install pywin32) — kontrola geometrii STP wyłączona")
    try:
        top = os.listdir(folder)
    except OSError as e:
        print("  BŁĄD dostępu:", e)
        return folder_niesprawdzony(folder, f"brak dostępu do folderu ({e})")
    xlsx = [f for f in top if f.lower().endswith((".xlsx", ".xlsm")) and not f.startswith("~$")]
    if not xlsx:
        print("  BŁĄD: brak pliku BOM (.xlsx) w folderze.")
        return folder_niesprawdzony(folder, "brak pliku BOM (.xlsx) w folderze")
    if len(xlsx) > 1:
        print(f"  UWAGA: kilka xlsx, biorę {xlsx[0]}")
    bom, err = czytaj_bom(os.path.join(folder, xlsx[0]))
    if err:
        print("  BŁĄD BOM:", err)
        return folder_niesprawdzony(folder, f"nie udało się odczytać BOM: {err}")

    bom_parts = {b["part"].upper() for b in bom}
    bom_asm = {b["part"].upper() for b in bom if "m" in b["flags"]}
    proj_root = re.sub(r"[\s\(\)].*$", "", os.path.basename(folder).upper())
    bom_asm.add(proj_root)

    pliki = {}
    pominiete = []
    for f in sorted(top):
        ext = os.path.splitext(f)[1].upper()
        if ext not in (".DXF", ".PDF"):
            continue
        if JUNK.search(f):
            pominiete.append(f)
            continue
        info = parsuj_nazwe(f, bom_parts, bom_asm)
        if info:
            info["file"] = f
            pliki.setdefault(info["part"].upper(), {})[ext] = info

    # --- geometria STP (raz na cały folder, nie per-część)
    stp_mapa, stp_blad = {}, None
    sciezka_stp = None if BEZ_STP else znajdz_stp(folder)
    uwagi_wstepne = []
    if sciezka_stp:
        print(f"  Otwieranie STP przez Inventor (może to potrwać kilkadziesiąt sekund): "
              f"{os.path.basename(sciezka_stp)}")
        stp_mapa, stp_blad = stp_dims_mapa(sciezka_stp)
        if stp_blad:
            uwagi_wstepne.append(f"STP ({os.path.basename(sciezka_stp)}): {stp_blad}")
        elif stp_mapa:
            print(f"  STP: odczytano geometrię {len(stp_mapa)} części.")
    else:
        uwagi_wstepne.append("brak pliku .stp/.step w folderze — kontrola geometrii 3D pominięta "
                              "(sam brak pliku nie jest błędem, nie wszystkie projekty go mają)")

    bledy, uwagi, bez_stempla, weryfikacje = [], [], [], []
    stp_info = []  # różnice z modelu .stp — tylko informacyjnie, nie są błędami
    uwagi.extend(uwagi_wstepne)
    ok = 0
    dxf_sprawdzone, dxf_bledne, dxf_uwagi = 0, [], 0

    for b in bom:
        uwagi_przed = len(uwagi)
        kat = kategoria(b, proj_root)
        f = pliki.pop(b["part"].upper(), {})
        dxf, pdf = f.get(".DXF"), f.get(".PDF")
        problemy = []
        etykieta = f"[{b['pos']}] {b['part']} ({b['title'] or b['desc'][:25] or kat})"
        weryf = None
        if dxf and not dxf["asm"]:
            weryf = dict(folder=os.path.basename(folder), pos=b["pos"], part=b["part"], tytul=b["title"],
                         plik=dxf["file"], potwierdzony=False, nakladka=None, kontur_ok=None, powod="")
            if not GEO:
                weryf["powod"] = "brak biblioteki ezdxf — geometria DXF nie jest sprawdzana"
            elif not AUTOTEST["ok"]:
                weryf["powod"] = f"AUTOTEST geometrii nie przeszedł na tym komputerze ({AUTOTEST['opis']})"
        try:

            # --- wymagane pliki wg kategorii
            if kat == "purchased":
                if not dxf and not pdf:
                    continue  # norma (93% w archiwum) — bez szumu
            elif kat == "ramka":
                if not pdf:
                    uwagi.append(f"{etykieta}: ramka tabliczki bez PDF — w archiwum 61% ma PDF, 20% nie; do decyzji")
                    continue
            elif kat == "zlozenie":
                if not pdf:
                    problemy.append("BRAK PDF złożenia")
                if dxf:
                    uwagi.append(f"{etykieta}: złożenie ma DXF — rzadkie (szablony G BOLTON), sprawdź czy celowe")
            elif kat == "blacha":
                if not dxf:
                    if czy_krawedz_plaskownik(b):
                        uwagi.append(f"{etykieta}: krawędź z paska {b['desc']} bez DXF — typowe dla GRADING "
                                     f"(płaskownik, cięcie+ukos), ale flaga C w BOM by to potwierdziła (jest: "
                                     f"'{''.join(sorted(b['flags'])) or '-'}')")
                    else:
                        problemy.append("BRAK DXF (blacha wypalana — 99% archiwum ma DXF)")
                if not pdf:
                    problemy.append("BRAK PDF")
            elif kat == "blacha_c":
                if not pdf:
                    problemy.append("BRAK PDF")
            elif kat == "przygotowka":
                if not dxf:
                    problemy.append("BRAK DXF (przygotówka/PL-t)")
            else:  # material_ciety, inne
                if not pdf:
                    if kat == "inne" and not b["desc"] and not dxf:
                        uwagi.append(f"{etykieta}: brak plików, puste DESCRIPTION — wygląda na część kupną "
                                     f"oznaczoną Normal; sprawdź BOM STRUCTURE")
                        continue
                    problemy.append("BRAK PDF")

            # --- nazwa pliku vs BOM
            for info in (dxf, pdf):
                if not info:
                    continue
                e = info["ext"]
                if info["asm"]:
                    if info["flags"] != b["flags"]:
                        problemy.append(f"{e}: flagi '{''.join(sorted(info['flags'])) or '-'}' vs BOM "
                                        f"'{''.join(sorted(b['flags'])) or '-'}'")
                    continue
                if info["thk"] is not None and b["dims"]:
                    if info["thk"] not in b["dims"]:
                        problemy.append(f"{e}: grubość {info['thk']:g} nie występuje w wymiarach BOM ({b['desc']})")
                if kat == "blacha" and info["thk"] is None:
                    problemy.append(f"{e}: brak grubości w nazwie ({b['desc']})")
                if info["mat"] and not mat_zgodny(info["mat"], b["mat"]):
                    problemy.append(f"{e}: materiał '{info['mat']}' vs BOM '{b['mat']}'")
                if info["qty"] is not None and b["qty"] is not None:
                    if info["qty"] not in (b["qty"], b["qty_total"]):
                        problemy.append(f"{e}: ilość {info['qty']} vs BOM {b['qty']}"
                                        + (f" (całk. {b['qty_total']:g})" if b["qty_total"] != b["qty"] else ""))
                if info["flags"] != b["flags"]:
                    problemy.append(f"{e}: flagi '{''.join(sorted(info['flags'])) or '-'}' vs BOM "
                                    f"'{''.join(sorted(b['flags'])) or '-'}'")

            # --- prefiksy
            if dxf and not dxf["prefiks"]:
                uwagi.append(f"{etykieta}: DXF bez prefiksu złożenia w nazwie")
            if pdf and pdf["prefiks"] and not pdf["asm"]:
                uwagi.append(f"{etykieta}: PDF z prefiksem złożenia (wg wzorca PDF-y części są bez)")

            # --- zawartosc PDF
            tresc = None
            if pdf:
                tresc = czytaj_pdf(os.path.join(folder, pdf["file"]))
                if "stamp" not in tresc:
                    bez_stempla.append(etykieta)
                elif b["qty"] is not None and tresc["stamp"] not in (b["qty"], b["qty_total"]):
                    if tresc.get("lustro"):
                        uwagi.append(f"{etykieta}: stempel 'Ilość: {tresc['stamp']}' vs BOM {b['qty']}, ale jest "
                                     f"adnotacja o odbiciu lustrzanym — prawdopodobnie celowe")
                    else:
                        problemy.append(f"stempel 'Ilość: {tresc['stamp']}' vs BOM {b['qty']}")
                if kat in ("blacha", "blacha_c") and "dims" in tresc and b["dims"]:
                    if tresc["dims"] != sorted(b["dims"]):
                        problemy.append(f"tabliczka PDF: wymiary {tresc['dims']} vs BOM {sorted(b['dims'])}")
                if kat in ("blacha", "blacha_c") and "mat" in tresc and not mat_zgodny(tresc["mat"], b["mat"]):
                    problemy.append(f"tabliczka PDF: materiał '{tresc['mat']}' vs BOM '{b['mat']}'")

            # --- geometria DXF vs DESCRIPTION (tylko wersja GEO)
            if GEO and dxf and not dxf["asm"] and len(b["dims"]) == 3:
                bb = None
                try:
                    bb = dxf_gabaryt(os.path.join(folder, dxf["file"]))
                except Exception as e:
                    uwagi.append(f"{etykieta}: nie udało się odczytać geometrii DXF ({type(e).__name__})")
                if bb:
                    dev, plan, bbs = geo_porownaj(b, dxf, bb)
                    opis_geo = (f"geometria DXF {bbs[0]:.0f}x{bbs[1]:.0f} vs BOM {plan[0]:g}x{plan[1]:g} "
                                f"(odchyłka {dev:.1f} mm)")
                    gieta = "p" in b["flags"]
                    obrabiana = "o" in b["flags"] or "u" in b["flags"]
                    if max(plan) > 0 and max(bbs) > 1.5 * max(plan):
                        uwagi.append(f"{etykieta}: DXF zawiera więcej niż kontur części "
                                     f"(ramka/arkusz? {opis_geo})")
                    elif gieta:
                        if dev > GEO_TOL_BLAD:
                            uwagi.append(f"{etykieta}: {opis_geo} — część gięta/walcowana; jeśli DESCRIPTION "
                                         f"to wymiar po gięciu, różnica jest naturalna")
                    elif obrabiana:
                        if dev > GEO_TOL_BLAD:
                            uwagi.append(f"{etykieta}: {opis_geo} — część z obróbką/ukosem (możliwy naddatek); "
                                         f"zweryfikuj")
                    else:
                        if dev > GEO_TOL_BLAD:
                            problemy.append(opis_geo)
                        elif dev > GEO_TOL_INFO:
                            uwagi.append(f"{etykieta}: {opis_geo} — strefa szara")

            # --- gabaryt DXF vs tabliczka PDF, gdy BOM nie ma 3 wymiarów (np. "PL 10" przygotówki)
            if GEO and dxf and not dxf["asm"] and len(b["dims"]) != 3 and tresc and "dims" in tresc:
                try:
                    bb = dxf_gabaryt(os.path.join(folder, dxf["file"]))
                except Exception:
                    bb = None
                if bb:
                    dev, plan, bbs = geo_porownaj(dict(b, dims=tresc["dims"]), dxf, bb)
                    opis_geo = (f"geometria DXF {bbs[0]:.0f}x{bbs[1]:.0f} vs rysunek PDF {plan[0]:g}x{plan[1]:g} "
                                f"(odchyłka {dev:.1f} mm)")
                    if dev > GEO_TOL_BLAD and not ("p" in b["flags"] or "o" in b["flags"] or "u" in b["flags"]):
                        problemy.append(opis_geo)
                    elif dev > GEO_TOL_INFO:
                        uwagi.append(f"{etykieta}: {opis_geo} — strefa szara / część gięta lub obrabiana")

            # --- nakładka DXF na widok z rysunku PDF (1:1)
            nak_stan = None
            if weryf is not None and GEO and pdf:
                nak = None
                try:
                    nak = nakladka(os.path.join(folder, dxf["file"]), os.path.join(folder, pdf["file"]))
                except Exception as e:
                    weryf["powod"] = f"błąd programu przy nakładce ({type(e).__name__}: {e})"
                if nak:
                    bl_n, uw_n, potwierdzony = ocen_nakladke(b, nak)
                    problemy.extend(bl_n)
                    uwagi.extend(f"{etykieta}: {u}" for u in uw_n)
                    if nak["mediana"] <= NAKLADKA_ZLY * nak["tol"]:
                        nak_stan = "potwierdzony" if potwierdzony else "znaleziony"
                    weryf.update(nakladka=nak, potwierdzony=potwierdzony and not nak["lustro"])
                    if nak["lustro"]:
                        weryf["powod"] = "pasuje do rysunku tylko po odbiciu lustrzanym — sprawdź stronę gięcia/ukosu"
                    elif uw_n and not potwierdzony:
                        weryf["powod"] = "; ".join(u.split(" — ", 1)[-1] for u in uw_n)
                    if bl_n or (uw_n and nak_stan) or PODGLAD_WSZYSTKIE:   # podgląd tylko przy znalezionym widoku
                        NAKLADKI_PDF.append((f"{os.path.basename(folder)} [{b['pos']}] {b['part']}",
                                             os.path.join(folder, pdf["file"]), nak["strona"], nak["Q"], nak["r"],
                                             nak["tol"]))
                elif not weryf["powod"]:
                    weryf["powod"] = ("na rysunku nie znaleziono widoku pasującego do DXF (brak widoku płaskiego/"
                                      "rozwinięcia, rysunek skanowany albo DXF zupełnie inny niż rysunek)")
            elif weryf is not None and GEO:
                weryf["powod"] = "brak PDF rysunku — nie ma z czym porównać DXF"

            # --- kontur DXF tak, jak zobaczy go wypalarka + fazy/otwory vs rysunek PDF
            if GEO and dxf and not dxf["asm"]:
                ana = None
                try:
                    ana = dxf_analiza(os.path.join(folder, dxf["file"]))
                except Exception as e:
                    problemy.append(f".DXF: kontur — nie udało się odczytać DXF ({type(e).__name__}: {e}) — plik "
                                    f"uszkodzony albo w nieobsługiwanym formacie")
                if ana:
                    bl_k, uw_k = ocen_kontur(b, kat, ana, tresc, nak_stan, weryf.get("nakladka"))
                    problemy.extend(bl_k)
                    uwagi.extend(f"{etykieta}: {u}" for u in uw_k)
                    weryf["kontur_ok"] = not bl_k

            # --- geometria STP (bryła 3D) vs DESCRIPTION — niezależne od DXF, tylko jeśli mamy mapę.
            # Tylko INFORMACYJNIE (poziom "INFO STP", nie liczy się do błędów): gabaryt bryły z STP
            # zależy od ułożenia części w złożeniu (obrót, pochylenie, gięcie), więc różnica przy
            # zgodnym DXF/PDF nie oznacza błędu dokumentacji (GE98022-S2, 2026-09-30).
            if stp_mapa and not (dxf and dxf.get("asm")) and len(b["dims"]) == 3:
                dims3 = stp_szukaj(stp_mapa, b["part"])
                if dims3:
                    wynik = geo_porownaj_3d(b["dims"], dims3)
                    if wynik:
                        dev, plan, znaleziony = wynik
                        opis_geo = (f".STP: bryła {znaleziony[0]:.0f}x{znaleziony[1]:.0f}x{znaleziony[2]:.0f} "
                                    f"vs BOM {plan[0]:g}x{plan[1]:g}x{plan[2]:g} (odchyłka {dev:.1f} mm)")
                        gieta = "p" in b["flags"]
                        obrabiana = "o" in b["flags"] or "u" in b["flags"]
                        if dev > GEO_TOL_BLAD:
                            if gieta:
                                powod = "część gięta — bryła w STP jest już po gięciu"
                            elif obrabiana:
                                powod = "część z obróbką — możliwy naddatek"
                            else:
                                powod = "najczęściej część obrócona/pochylona w złożeniu"
                            stp_info.append(f"{etykieta}: {opis_geo} — {powod}")
                elif b["part"].upper() != proj_root and "m" not in b["flags"] and kat not in ("purchased", "ramka"):
                    stp_info.append(f"{etykieta}: .STP: część nie znaleziona w bryle zespołu STP po nazwie "
                                    f"(inna nazwa w modelu albo część spoza tego pliku STP)")
        except Exception as e:
            problemy.append(f"BŁĄD PROGRAMU przy sprawdzaniu tej pozycji ({type(e).__name__}: {e}) — pozycja nie "
                            f"została sprawdzona do końca, sprawdź ją ręcznie")
            if weryf is not None:
                weryf["potwierdzony"] = False
                weryf["powod"] = f"błąd programu ({type(e).__name__}: {e})"

        if weryf is not None:
            geo_bl = [x for x in problemy if x.startswith(".DXF:") or x.startswith("geometria DXF")
                      or x.startswith("BŁĄD PROGRAMU")]
            if geo_bl and not all(x.startswith("BŁĄD PROGRAMU") for x in geo_bl):
                weryf["status"] = "RÓŻNICE"
                weryf["powod"] = "; ".join(x.replace(".DXF: ", "") for x in geo_bl)
            elif weryf["potwierdzony"] and weryf["kontur_ok"] and AUTOTEST["ok"]:
                weryf["status"] = "ZGODNY 1:1"
            else:
                weryf["status"] = "NIEZWERYFIKOWANY"
                if not weryf["powod"]:
                    weryf["powod"] = "nie udało się potwierdzić konturu (szczegóły w zakładce Błędy i uwagi)"
            WERYFIKACJA_DXF.append(weryf)
            weryfikacje.append(weryf)

        # --- zgodność DXF tej pozycji (do podsumowania na końcu raportu)
        if dxf and not dxf["asm"]:
            dxf_sprawdzone += 1
        if any("DXF" in p for p in problemy):
            dxf_bledne.append(b["part"])
        dxf_uwagi += sum(1 for u in uwagi[uwagi_przed:] if "DXF" in u)

        if problemy:
            bledy.append(f"{etykieta}: " + "; ".join(problemy))
        else:
            ok += 1

    if bez_stempla:
        if len(bez_stempla) > 5:
            uwagi.append(f"{len(bez_stempla)} PDF-ów bez stempla 'Ilość:' (folder nie przeszedł przez "
                         f"PdfRenamer / starsze wydanie?), np. {', '.join(bez_stempla[:3])}")
        else:
            for e in bez_stempla:
                uwagi.append(f"{e}: PDF bez stempla 'Ilość:' (nie przeszedł przez PdfRenamer?)")

    for key, f in pliki.items():
        if key == proj_root or key in bom_asm:
            continue  # glowne zlozenie bez wlasnego wiersza w BOM — norma w starszych wydaniach
        uwagi.append(f"{key}: plik(i) {', '.join(v['file'] for v in f.values())} bez pozycji w BOM")
        if ".DXF" in uwagi[-1].upper():
            dxf_uwagi += 1

    print(f"  BOM: {len(bom)} pozycji | zgodnych: {ok} | błędów: {len(bledy)} | uwag: {len(uwagi)}"
          + (f" | info STP: {len(stp_info)}" if stp_info else "")
          + (f" | pominięte wytyczne: {len(pominiete)}" if pominiete else ""))
    if bledy:
        print("  --- BŁĘDY " + "-" * 62)
        for x in bledy:
            print("   !!", x)
    if uwagi:
        print("  --- UWAGI " + "-" * 62)
        for x in uwagi:
            print("    -", x)
    if stp_info:
        print(f"  --- INFO STP: {len(stp_info)} porównań z modelem 3D .stp — tylko informacyjnie, to NIE są błędy "
              f"(pełna lista w raporcie, zakładka 'Model STP (info)'), np.:")
        for x in stp_info[:3]:
            print("    i", x)
    if dxf_sprawdzone and not dxf_bledne:
        print(f"  >>> WSZYSTKIE PLIKI DXF ZGODNE Z BOM ({dxf_sprawdzone}) <<<")
    zgodne_1do1 = [w for w in weryfikacje if w["status"] == "ZGODNY 1:1"]
    niezweryf = [w for w in weryfikacje if w["status"] == "NIEZWERYFIKOWANY"]
    if weryfikacje:
        maks = max((w["nakladka"]["r"].max() for w in zgodne_1do1), default=0)
        print(f"  >>> DXF POTWIERDZONE 1:1 Z RYSUNKIEM: {len(zgodne_1do1)}/{len(weryfikacje)}"
              + (f" (maks. odchyłka {maks:.2f} mm)" if zgodne_1do1 else "") + " <<<")
        for w in niezweryf:
            print(f"    ?? [{w['pos']}] {w['part']}: NIEZWERYFIKOWANY — {w['powod']}")
    if not bledy and not niezweryf:
        print("  >>> FOLDER ZGODNY Z WZORCEM <<<")
    elif not bledy:
        print(f"  >>> BRAK BŁĘDÓW, ALE {len(niezweryf)} DXF NIEZWERYFIKOWANYCH — SPRAWDŹ JE RĘCZNIE <<<")

    dodaj_wiersze(folder, bledy, "BŁĄD")
    dodaj_wiersze(folder, uwagi, "UWAGA")
    dodaj_wiersze(folder, stp_info, "INFO STP")
    PODSUMOWANIE_FOLDEROW.append(dict(
        folder=os.path.basename(folder), sciezka=folder,
        ok=ok, bledy=len(bledy), uwagi=len(uwagi), pozycji=len(bom),
        stp="TAK" if stp_mapa else ("BŁĄD" if stp_blad else "brak pliku"),
        stp_info=len(stp_info),
        dxf_sprawdzone=dxf_sprawdzone, dxf_bledne=dxf_bledne, dxf_uwagi=dxf_uwagi,
        dxf_1do1=len(zgodne_1do1), dxf_wszystkie=len(weryfikacje),
        dxf_niezweryf=[(w["part"], w["powod"]) for w in niezweryf],
        dxf_rozne=[w["part"] for w in weryfikacje if w["status"] == "RÓŻNICE"],
    ))
    return 1 if bledy or (niezweryf and NIEZWERYFIKOWANE_BLOKUJA) else 0


# ---------------------------------------------------------------- raport xlsx
def zapisz_raport_xlsx(sciezka):
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Błędy i uwagi"
    naglowki = ["Folder", "POS (miejsce w BOM)", "Część", "Tytuł/opis", "Poziom",
                "Plik", "Rodzaj sprawdzenia", "Opis problemu"]
    ws.append(naglowki)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="404040")
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

    czerwony = PatternFill("solid", fgColor="F8CBCB")
    pomaranczowy = PatternFill("solid", fgColor="F8CBAD")
    zolty = PatternFill("solid", fgColor="FFF3B0")
    niebieski = PatternFill("solid", fgColor="DDEBF7")
    zielony = PatternFill("solid", fgColor="C6EFCE")
    szary = PatternFill("solid", fgColor="EDEDED")
    kolor_poziomu = {"BŁĄD": czerwony, "UWAGA": zolty, "INFO STP": niebieski}
    kolejnosc_poziomu = {"BŁĄD": 0, "UWAGA": 1, "INFO STP": 2}

    # Wiersze z modelu .stp są tylko informacyjne (gabaryt bryły zależy od ułożenia w złożeniu) — w głównej
    # zakładce zostają tylko przy częściach, które mają też błąd/uwagę z DXF/PDF (jak radzi PORADNIK);
    # reszta idzie do osobnej zakładki, żeby nie zasłaniała prawdziwych błędów.
    z_problemem = {(w["folder"], w["part"]) for w in WSZYSTKIE_WIERSZE if w["poziom"] in ("BŁĄD", "UWAGA")}
    stp_osobno = [w for w in WSZYSTKIE_WIERSZE if w["poziom"] == "INFO STP" and (w["folder"], w["part"]) not in z_problemem]
    wiersze = sorted([w for w in WSZYSTKIE_WIERSZE if not (w["poziom"] == "INFO STP"
                                                           and (w["folder"], w["part"]) not in z_problemem)],
                     key=lambda w: (w["folder"], kolejnosc_poziomu.get(w["poziom"], 1), pos_key(w["pos"])))
    for w in wiersze:
        ws.append([w["folder"], w["pos"], w["part"], w["tytul"], w["poziom"],
                   w["plik"], w["rodzaj"], w["opis"]])
        fill = kolor_poziomu.get(w["poziom"], zolty)
        for c in ws[ws.max_row]:
            c.fill = fill
            c.alignment = Alignment(vertical="top", wrap_text=(c.column_letter == "H"))

    szerokosci = [24, 14, 16, 26, 10, 6, 20, 70]
    for i, sz in enumerate(szerokosci, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = sz
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:H{ws.max_row}"

    # --- podsumowanie na końcu arkusza (poza autofiltrem): DXF, model .stp, werdykt
    ws.append([])
    ws.append(["PODSUMOWANIE"])
    ws[ws.max_row][0].font = Font(bold=True)

    def wiersz_podsumowania(folder, etykieta, tekst, fill):
        ws.append([folder, etykieta, tekst])
        r = ws.max_row
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
        for c in ws[r]:
            c.fill = fill
            c.alignment = Alignment(vertical="top", wrap_text=True)
        ws[r][1].font = Font(bold=True)
        ws.row_dimensions[r].height = 30

    for p in PODSUMOWANIE_FOLDEROW:
        if p.get("niesprawdzony"):
            wiersz_podsumowania(p["folder"], "Wynik", f"FOLDER NIE ZOSTAŁ SPRAWDZONY: {p['niesprawdzony']}", czerwony)
            continue
        n, zle, uw = p["dxf_sprawdzone"], p["dxf_bledne"], p["dxf_uwagi"]
        if zle:
            wiersz_podsumowania(p["folder"], "DXF", f"DXF NIEZGODNE Z BOM: {len(zle)} poz. "
                                f"({', '.join(zle)}) — czerwone wiersze wyżej.", czerwony)
        elif n:
            wiersz_podsumowania(p["folder"], "DXF", f"Wszystkie pliki DXF zgodne z BOM ({n} szt.: grubość, "
                                f"materiał, ilość, flagi, geometria)."
                                + (f" Uwagi do DXF wyżej ({uw}) są do przejrzenia, to nie są błędy." if uw else ""),
                                zielony)
        else:
            wiersz_podsumowania(p["folder"], "DXF", "Brak plików DXF do sprawdzenia.", szary)

        if p["stp_info"]:
            wiersz_podsumowania(p["folder"], "Model .stp",
                                f"Porównania z modelem 3D (.stp): {p['stp_info']} — są w zakładce 'Model STP (info)'; "
                                f"tu widać je tylko przy częściach, które mają też błąd/uwagę z DXF/PDF. Gabaryt bryły "
                                f"w .stp zależy od ułożenia części w złożeniu (obrót, pochylenie, gięcie), więc taka "
                                f"różnica NIE oznacza błędu dokumentacji, jeśli DXF i PDF są zgodne z BOM.", niebieski)
        elif p["stp"] == "TAK":
            wiersz_podsumowania(p["folder"], "Model .stp", "Model 3D (.stp) zgodny z BOM.", zielony)

        if p["dxf_wszystkie"]:
            if p["dxf_rozne"]:
                wiersz_podsumowania(p["folder"], "DXF 1:1", f"DXF RÓŻNE OD RYSUNKU: {len(p['dxf_rozne'])} szt. "
                                    f"({', '.join(p['dxf_rozne'])}) — szczegóły w zakładce Weryfikacja DXF 1 do 1 i w "
                                    f"podglądzie NAKLADKI_*.pdf.", czerwony)
            if p["dxf_niezweryf"]:
                wiersz_podsumowania(p["folder"], "DXF 1:1", f"{len(p['dxf_niezweryf'])} z {p['dxf_wszystkie']} DXF "
                                    f"NIEZWERYFIKOWANYCH — sprawdź je ręcznie: "
                                    + "; ".join(f"{c} ({pw})" for c, pw in p["dxf_niezweryf"][:5])
                                    + (" …" if len(p["dxf_niezweryf"]) > 5 else ""), pomaranczowy)
            if not p["dxf_rozne"] and not p["dxf_niezweryf"]:
                wiersz_podsumowania(p["folder"], "DXF 1:1", f"Wszystkie {p['dxf_wszystkie']} DXF potwierdzone 1:1 z "
                                    f"rysunkami (nakładka konturu na widok + czysty kontur).", zielony)

        if p["bledy"]:
            wiersz_podsumowania(p["folder"], "Wynik", f"Błędów w dokumentacji: {p['bledy']} — czerwone wiersze wyżej.",
                                czerwony)
        elif p["dxf_niezweryf"]:
            wiersz_podsumowania(p["folder"], "Wynik", f"Brak błędów, ale {len(p['dxf_niezweryf'])} DXF "
                                f"NIEZWERYFIKOWANYCH — dokumentacja NIE jest w pełni potwierdzona, sprawdź je ręcznie.",
                                pomaranczowy)
        else:
            wiersz_podsumowania(p["folder"], "Wynik", "Dokumentacja zgodna z BOM — brak błędów."
                                + (" Wszystkie DXF potwierdzone 1:1 z rysunkami." if p["dxf_wszystkie"] else "")
                                + (f" Uwagi ({p['uwagi']}) są do przejrzenia." if p["uwagi"] else ""), zielony)

    # --- zakładka: informacje z modelu .stp (bez związku z błędami DXF/PDF)
    ws4 = wb.create_sheet("Model STP (info)")
    ws4.append(["Folder", "POS", "Część", "Tytuł/opis", "Opis (tylko informacyjnie — to nie są błędy)"])
    for c in ws4[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="404040")
    for w in sorted(stp_osobno, key=lambda w: (w["folder"], pos_key(w["pos"]))):
        ws4.append([w["folder"], w["pos"], w["part"], w["tytul"], w["opis"]])
        for c in ws4[ws4.max_row]:
            c.fill = niebieski
            c.alignment = Alignment(vertical="top", wrap_text=(c.column_letter == "E"))
    for i, sz in enumerate([22, 8, 14, 24, 100], start=1):
        ws4.column_dimensions[openpyxl.utils.get_column_letter(i)].width = sz
    if ws4.max_row > 1:
        ws4.auto_filter.ref = f"A1:E{ws4.max_row}"

    # --- zakładka: każdy DXF z dowodem weryfikacji 1:1
    ws3 = wb.create_sheet("Weryfikacja DXF 1 do 1")   # w nazwie arkusza nie wolno ":"
    ws3.append(["Folder", "POS", "Część", "Tytuł", "Plik DXF", "Status", "Maks. odchyłka [mm]",
                "Tolerancja [mm]", "Widok na rysunku", "Orientacja DXF", "Powód / szczegóły"])
    for c in ws3[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="404040")
    ws3.freeze_panes = "A2"
    kolor_statusu = {"ZGODNY 1:1": zielony, "RÓŻNICE": czerwony, "NIEZWERYFIKOWANY": pomaranczowy}
    for w in sorted(WERYFIKACJA_DXF, key=lambda w: (w["folder"], pos_key(w["pos"]))):
        n = w.get("nakladka")
        ws3.append([w["folder"], w["pos"], w["part"], w["tytul"], w["plik"], w["status"],
                    round(float(n["r"].max()), 2) if n else None, round(n["tol"], 2) if n else None,
                    (f"strona {n['strona'] + 1}, 1:{_mm(n['skala'])}" + ("" if n["skala_opisana"] else " (dobrana)")
                     + (", szukany na całej stronie" if n.get("zrodlo") == "strona" else "")) if n else "",
                    (f"obrót {n['obrot']:g}°" + (", LUSTRO" if n["lustro"] else "")) if n else "",
                    w["powod"]])
        for c in ws3[ws3.max_row]:
            c.alignment = Alignment(vertical="top", wrap_text=(c.column_letter in "EK"))
        ws3[ws3.max_row][5].fill = kolor_statusu.get(w["status"], szary)
    for i, sz in enumerate([22, 8, 14, 22, 40, 18, 12, 11, 24, 16, 70], start=1):
        ws3.column_dimensions[openpyxl.utils.get_column_letter(i)].width = sz
    if ws3.max_row > 1:
        ws3.auto_filter.ref = f"A1:K{ws3.max_row}"

    ws2 = wb.create_sheet("Podsumowanie folderów")
    ws2.append(["Folder", "Pozycji BOM", "Zgodnych", "Błędów", "Uwag", "DXF", "Plik STP", "Info STP",
                "DXF 1:1", "Niezweryf."])
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="404040")
    for p in PODSUMOWANIE_FOLDEROW:
        if p["dxf_bledne"]:
            dxf_stan = f"NIEZGODNE ({len(p['dxf_bledne'])})"
        elif p["dxf_sprawdzone"]:
            dxf_stan = f"ZGODNE ({p['dxf_sprawdzone']})"
        else:
            dxf_stan = "brak"
        ws2.append([p["folder"], p["pozycji"], p["ok"], p["bledy"], p["uwagi"], dxf_stan, p["stp"], p["stp_info"],
                    f"{p['dxf_1do1']}/{p['dxf_wszystkie']}", len(p["dxf_niezweryf"])])
        if p["bledy"] > 0:
            for c in ws2[ws2.max_row]:
                c.fill = czerwony
        else:
            ws2[ws2.max_row][5].fill = zielony if p["dxf_sprawdzone"] else szary
        if p["stp_info"]:
            ws2[ws2.max_row][7].fill = niebieski
        if p["dxf_wszystkie"]:
            ws2[ws2.max_row][8].fill = (czerwony if p["dxf_rozne"] else pomaranczowy if p["dxf_niezweryf"]
                                        else zielony)
            if p["dxf_niezweryf"]:
                ws2[ws2.max_row][9].fill = pomaranczowy
    for i, sz in enumerate([28, 12, 10, 8, 8, 14, 10, 9, 9, 10], start=1):
        ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = sz
    if ws2.max_row > 1:
        ws2.auto_filter.ref = f"A1:J{ws2.max_row}"

    wb.save(sciezka)
    return sciezka


WERSJA = "v5 (2026-10-08: nakładka DXF 1:1, statusy weryfikacji, autotest)"


class _Dziennik:
    """Wszystko, co program wypisuje w konsoli, trafia też do pliku LOG_SPRAWDZENIA_<data>.txt obok raportu
    — gdy coś pójdzie nie tak, wystarczy przysłać ten plik."""

    def __init__(self, strumien, sciezka):
        self.strumien, self.plik = strumien, None
        try:
            self.plik = open(sciezka, "w", encoding="utf-8")
        except OSError:
            pass

    def write(self, t):
        self.strumien.write(t)
        if self.plik:
            self.plik.write(t)
            self.plik.flush()
        return len(t)

    def flush(self):
        self.strumien.flush()


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(0)
    kod = 0
    folder_dla_raportu = None
    znacznik = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    foldery = [a for a in sys.argv[1:] if os.path.isdir(a)]
    if foldery:
        sys.stdout = _Dziennik(sys.stdout, os.path.join(os.path.dirname(foldery[0].rstrip("\\/")) or foldery[0],
                                                         f"LOG_SPRAWDZENIA_{znacznik}.txt"))
    print(f"sprawdz_hybrydowo_GEO {WERSJA} — folderów do sprawdzenia: {len(foldery)}")
    if GEO:
        AUTOTEST.update(autotest())
        print(("Autotest geometrii DXF: OK — " if AUTOTEST["ok"] else
               "\n" + "!" * 92 + "\n!!! AUTOTEST GEOMETRII NIE PRZESZEDŁ — wyniki DXF NIE SĄ WIARYGODNE: ")
              + AUTOTEST["opis"] + ("" if AUTOTEST["ok"] else "\n" + "!" * 92))
        if not AUTOTEST["ok"]:
            kod = 1
            WSZYSTKIE_WIERSZE.append(dict(folder="(program)", pos="", part="AUTOTEST", tytul="", poziom="BŁĄD",
                                          plik="", rodzaj="autotest programu",
                                          opis=f"AUTOTEST geometrii nie przeszedł — żaden DXF nie jest uznany "
                                               f"za zweryfikowany: {AUTOTEST['opis']}"))
    else:
        print("UWAGA: brak biblioteki ezdxf — DXF NIE będą weryfikowane (wszystkie dostaną status NIEZWERYFIKOWANY)")
    for arg in sys.argv[1:]:
        if os.path.isdir(arg):
            if folder_dla_raportu is None:
                folder_dla_raportu = os.path.dirname(arg.rstrip("\\/")) or arg
            try:
                kod |= sprawdz(arg)
            except Exception as e:
                import traceback
                traceback.print_exc()
                print(f"  BŁĄD PROGRAMU — folder NIE został sprawdzony: {type(e).__name__}: {e}")
                kod |= folder_niesprawdzony(arg, f"błąd programu {type(e).__name__}: {e}")
        else:
            print("Pominięto (to nie folder):", arg)

    if WSZYSTKIE_WIERSZE or PODSUMOWANIE_FOLDEROW:
        sciezka_raport = os.path.join(folder_dla_raportu or ".", f"RAPORT_SPRAWDZENIA_{znacznik}.xlsx")
        try:
            zapisz_raport_xlsx(sciezka_raport)
            print(f"\nZapisano raport zbiorczy: {sciezka_raport}")
            if GEO and NAKLADKI_PDF:
                sciezka_nak = zapisz_nakladki_pdf(os.path.join(folder_dla_raportu or ".", f"NAKLADKI_{znacznik}.pdf"))
                print(f"Podgląd nakładek DXF na rysunki (części z różnicami): {sciezka_nak}")
            if not os.environ.get("WYDAJ_FOLDER_PIPELINE"):
                try:
                    os.startfile(sciezka_raport)
                except Exception:
                    pass
        except Exception as e:
            print(f"\nNie udało się zapisać raportu xlsx: {e}")

    sys.exit(kod)
