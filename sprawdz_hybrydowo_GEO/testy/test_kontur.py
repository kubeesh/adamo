# -*- coding: utf-8 -*-
"""Testy kontroli konturu DXF / faz / otworów na sztucznych plikach.

Uruchom:  python testy/test_kontur.py      (albo: python -m pytest testy)
Wymaga:   pip install openpyxl pymupdf ezdxf
"""
import os
import sys
import tempfile

import ezdxf
import fitz
import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
import sprawdz_hybrydowo_GEO as S  # noqa: E402

TMP = tempfile.mkdtemp(prefix="test_kontur_")
WID = "Widoczne (ISO)"


def zapisz_dxf(nazwa, linie=(), okregi=(), polilinie=(), folder=TMP):
    """linie: (x1, y1, x2, y2[, warstwa]); okregi: (x, y, r[, warstwa]); polilinie: (pkt, warstwa)."""
    doc = ezdxf.new("R2010", units=ezdxf.units.MM)   # jak Inventor: $INSUNITS=4
    msp = doc.modelspace()
    for l in linie:
        msp.add_line(l[:2], l[2:4], dxfattribs={"layer": l[4] if len(l) > 4 else WID})
    for o in okregi:
        msp.add_circle(o[:2], o[2], dxfattribs={"layer": o[3] if len(o) > 3 else WID})
    for pkt, warstwa in polilinie:
        msp.add_lwpolyline(pkt, format="xyb", close=True, dxfattribs={"layer": warstwa})
    sciezka = os.path.join(folder, nazwa)
    doc.saveas(sciezka)
    return sciezka


def zapisz_pdf(nazwa, tekst, folder=TMP):
    doc = fitz.open()
    strona = doc.new_page()
    for i, linia in enumerate(tekst.split("\n")):
        strona.insert_text((50, 60 + 16 * i), linia)
    sciezka = os.path.join(folder, nazwa)
    doc.save(sciezka)
    return sciezka


def prostokat_z_fazami(f, w=200, h=100, warstwa=WID):
    """Prostokąt w x h z fazami f x f na 4 narożnikach, jako osobne LINE."""
    p = [(f, 0), (w - f, 0), (w, f), (w, h - f), (w - f, h), (f, h), (0, h - f), (0, f)]
    return [(*p[i], *p[(i + 1) % 8], warstwa) for i in range(8)]


def prostokat(w=200, h=100, warstwa=WID):
    p = [(0, 0), (w, 0), (w, h), (0, h)]
    return [(*p[i], *p[(i + 1) % 4], warstwa) for i in range(4)]


def bom(flags=""):
    return dict(pos="1", part="PG1", title="SIDE PLATE", desc="PL 10 x 200 x 100", dims=[10.0, 200.0, 100.0],
                mat="S355", flags=set(flags), qty=1, qty_total=1, struct="Normal")


def ocena(sciezka, tekst_pdf=None, flags="", kat="blacha"):
    ana = S.dxf_analiza(sciezka)
    tresc = S.wymiary_rysunku(tekst_pdf) if tekst_pdf is not None else None
    return ana, S.ocen_kontur(bom(flags), kat, ana, tresc)


# ------------------------------------------------------------------ testy

def test_tekst_rysunku():
    w = S.wymiary_rysunku("4x 10 x 45°\nn10,5\n2x Ø22\nM12\nR15\nPL 10 x 45 x 300\nFAZA 5x45\n3 x 30°")
    nogi = sorted(f["nogi"] for f in w["fazy"])
    assert (10, 10) in nogi and (5, 5) in nogi, nogi
    assert any(f["ile"] == 4 for f in w["fazy"])
    assert (1.73, 3) in nogi, nogi                      # 3 x 30° -> nogi 3 i 3*tg30
    assert not any(f["nogi"] == (10, 45) for f in w["fazy"])  # "PL 10 x 45 x 300" to nie faza
    assert 10.5 in w["fi"] and 22 in w["fi"], w["fi"]
    assert w["gwinty"] == [12] and w["promienie"] == [15]


def test_czysty_kontur_z_fazami():
    ana, (bl, uw) = ocena(zapisz_dxf("ok.dxf", prostokat_z_fazami(10)), "4x 10 x 45°")
    assert not ana["przerwy"] and not ana["wolne"] and not ana["rozgal"], ana
    assert len(ana["fazy"]) == 4 and all(f[:2] == (10, 10) for f in ana["fazy"]), ana["fazy"]
    assert bl == [] and uw == [], (bl, uw)


def test_zla_faza_vs_rysunek():
    _, (bl, uw) = ocena(zapisz_dxf("faza7.dxf", prostokat_z_fazami(7)), "4x 10 x 45°")
    # bez dopasowanego widoku (sam tekst) — UWAGA, bo notka może dotyczyć fazy krawędzi
    assert not bl and any("fazy — na rysunku faza 4x 10 x 45°" in x and "7x7 (4 szt.)" in x for x in uw), (bl, uw)


def test_zla_faza_przy_ukosowaniu_to_uwaga():
    _, (bl, uw) = ocena(zapisz_dxf("faza7u.dxf", prostokat_z_fazami(7)), "4x 10 x 45°", flags="u")
    assert not any("fazy" in x for x in bl) and any("fazy" in x for x in uw), (bl, uw)


def test_brak_fazy_w_dxf():
    _, (bl, uw) = ocena(zapisz_dxf("bez_faz.dxf", prostokat()), "10 x 45°")
    assert not bl and any("w DXF nie ma żadnej fazy" in x for x in uw), (bl, uw)


def test_za_malo_faz():
    # narożnik (200,0) bez fazy, pozostałe trzy z fazą 10x10
    linie = [(10, 0, 200, 0, WID), (200, 0, 200, 90, WID)] + prostokat_z_fazami(10)[3:]
    _, (bl, uw) = ocena(zapisz_dxf("3fazy.dxf", linie), "4x 10 x 45°")
    assert not bl and any("4x faza 10x10, w DXF tylko 3" in x for x in uw), (bl, uw)


def test_nieprzyciety_narozniki_po_fazie():
    # stary ostry narożnik został, faza dorysowana "na wierzch"
    ana, (bl, _) = ocena(zapisz_dxf("nieprzyciety.dxf", prostokat() + [(10, 0, 0, 10, "0")]), "")
    assert len(ana["rozgal"]) == 2, ana["rozgal"]
    assert any("rozgałęzień" in x for x in bl), bl


def test_przyciety_z_jednej_strony():
    linie = [(0, 0, 200, 0, WID), (200, 0, 200, 100, WID), (200, 100, 0, 100, WID), (0, 100, 0, 10, WID),
             (10, 0, 0, 10, "0")]
    ana, (bl, _) = ocena(zapisz_dxf("pol_przyciety.dxf", linie))
    assert len(ana["wolne"]) == 1 and len(ana["rozgal"]) == 1, ana
    assert any("wolnych końców" in x for x in bl) and any("rozgałęzień" in x for x in bl), bl


def test_przerwa_w_konturze():
    linie = prostokat()
    linie[0] = (0, 0, 199.5, 0, WID)
    ana, (bl, _) = ocena(zapisz_dxf("przerwa.dxf", linie))
    assert len(ana["przerwy"]) == 1 and abs(ana["przerwy"][0][0] - 0.5) < 0.01, ana
    assert any("przerw" in x and "0,5 mm" in x for x in bl), bl


def test_mikroprzerwa_ignorowana():
    linie = prostokat()
    linie[0] = (0, 0, 199.98, 0, WID)
    ana, _ = ocena(zapisz_dxf("mikro.dxf", linie))
    assert not ana["przerwy"] and not ana["wolne"], ana


def test_podwojny_kontur():
    ana, (bl, _) = ocena(zapisz_dxf("dubel.dxf", prostokat() + [(*l[2:4], *l[:2], WID) for l in prostokat()]))
    assert ana["duble"] == 4 and not ana["rozgal"], ana
    assert any("podwójne linie — 4" in x for x in bl), bl


def test_wspolsrodkowe_okregi():
    _, (bl, _) = ocena(zapisz_dxf("pogl.dxf", prostokat(), okregi=[(50, 50, 11), (50, 50, 13)]), "n22")
    assert any("współśrodkowe okręgi Ø22 / Ø26" in x for x in bl), bl


def test_dorysowany_otwor_bez_wymiaru():
    sciezka = zapisz_dxf("otwor0.dxf", prostokat(), okregi=[(50, 50, 6, "0")])
    _, (bl, uw) = ocena(sciezka, "PL 10 x 200 x 100")
    assert any("dorysowany otwór Ø12" in x for x in bl), bl
    assert any("dorysowane ręcznie" in x for x in uw), uw
    _, (bl, _) = ocena(sciezka, "n12")
    assert not any("otwór" in x for x in bl), bl


def test_otwor_z_eksportu_niezgodny_z_rysunkiem():
    _, (bl, uw) = ocena(zapisz_dxf("otwor_eksp.dxf", prostokat(), okregi=[(50, 50, 7)]), "2x n12")
    assert any("otwór Ø14 z DXF nie występuje" in x for x in bl), (bl, uw)
    _, (bl, uw) = ocena(zapisz_dxf("gwint.dxf", prostokat(), okregi=[(50, 50, 5.1)]), "M12")
    assert not bl and not uw, (bl, uw)   # otwór pod gwint M12 (Ø10,2)


def test_dorysowane_wystaje():
    _, (bl, _) = ocena(zapisz_dxf("wystaje.dxf", prostokat() + [(300, 0, 300, 50, "0")]))
    assert any("wystaje 100 mm" in x for x in bl), bl


def test_polilinia_z_lukami_i_linie_giecia():
    pkt = [(0, 0, 0), (190, 0, 0.4142), (200, 10, 0), (200, 100, 0), (0, 100, 0)]   # łuk R10 w narożniku
    sciezka = zapisz_dxf("poly.dxf", [(100, 0, 100, 100, "IV_BEND")], polilinie=[(pkt, "IV_OUTER_PROFILE")])
    ana, (bl, uw) = ocena(sciezka, "")
    assert not ana["przerwy"] and not ana["wolne"] and not ana["rozgal"] and not ana["fazy"], ana
    assert bl == [] and uw == [], (bl, uw)


def test_faza_niezwymiarowana_bez_notek():
    _, (bl, uw) = ocena(zapisz_dxf("faza15.dxf", prostokat_z_fazami(15)), "200\n100")
    assert not bl and any("faza 15x15 (4 szt.) z DXF nie jest zwymiarowana" in x for x in uw), (bl, uw)
    _, (_, uw) = ocena(zapisz_dxf("faza15b.dxf", prostokat_z_fazami(15)), "200\n100\n15")
    assert not uw, uw


def test_jednostki_cale():
    doc = ezdxf.new("R2010", units=ezdxf.units.IN)
    for l in prostokat():
        doc.modelspace().add_line(l[:2], l[2:4], dxfattribs={"layer": WID})
    sciezka = os.path.join(TMP, "cale.dxf")
    doc.saveas(sciezka)
    _, (bl, uw) = ocena(sciezka)
    assert not bl and any("calach" in x for x in uw), (bl, uw)


def test_blok_insert_i_okrag_odwrocony():
    """Geometria schowana w bloku (INSERT) i okrąg z wektorem wyciągnięcia (0,0,-1)."""
    doc = ezdxf.new("R2010", units=ezdxf.units.MM)
    blk = doc.blocks.new("CZESC")
    for l in prostokat():
        blk.add_line(l[:2], l[2:4], dxfattribs={"layer": "0"})
    msp = doc.modelspace()
    msp.add_blockref("CZESC", (0, 0), dxfattribs={"layer": WID})
    msp.add_circle((-50, 50), 5, dxfattribs={"layer": WID, "extrusion": (0, 0, -1)})   # w WCS: (50, 50)
    sciezka = os.path.join(TMP, "blok.dxf")
    doc.saveas(sciezka)
    ana = S.dxf_analiza(sciezka)
    assert not ana["przerwy"] and not ana["wolne"] and not ana["rozgal"] and not ana["pusty"], ana
    assert ana["reczne"] == 0, ana                                  # warstwa "0" w bloku dziedziczy warstwę wstawienia
    assert abs(ana["okregi"][0][1][0] - 50) < 1e-6, ana["okregi"]   # środek przeliczony do WCS


def test_pusty_dxf():
    sciezka = zapisz_dxf("pusty.dxf", [(0, 0, 100, 0, "Wymiar (ISO)")])
    _, (bl, _) = ocena(sciezka)
    assert any("nie ma żadnej geometrii" in x for x in bl), bl


def test_autotest_wykrywa_awarie():
    """Autotest musi przejść normalnie i NIE przejść, gdy pomiar geometrii jest zepsuty
    (np. inna wersja biblioteki na komputerze użytkownika)."""
    assert S.autotest()["ok"], S.autotest()
    oryginal = S._odl_do_odcinkow
    try:
        S._odl_do_odcinkow = lambda P, seg: (_ for _ in ()).throw(TypeError("symulowana awaria numpy"))
        assert not S.autotest()["ok"]
        S._odl_do_odcinkow = lambda P, seg: (oryginal(P, seg)[0] * 0, oryginal(P, seg)[1])   # "wszystko pasuje"
        assert not S.autotest()["ok"]
    finally:
        S._odl_do_odcinkow = oryginal


def test_stp_dopasowanie_nazw():
    mapa = {"PG4136348:1": (1, 2, 3), "PG999 - TOP PLATE": (4, 5, 6), "PG12.IPT": (7, 8, 9), "PG123": (0, 0, 0)}
    assert S.stp_szukaj(mapa, "PG4136348") == (1, 2, 3)
    assert S.stp_szukaj(mapa, "pg999") == (4, 5, 6)
    assert S.stp_szukaj(mapa, "PG12") == (7, 8, 9)        # nie myli z PG123
    assert S.stp_szukaj(mapa, "PG1") is None


def test_raport_stp_w_osobnej_zakladce():
    S.WSZYSTKIE_WIERSZE.clear()
    S.PODSUMOWANIE_FOLDEROW.clear()
    S.WERYFIKACJA_DXF.clear()
    w = dict(folder="F", pos="1", tytul="", plik="", rodzaj="")
    S.WSZYSTKIE_WIERSZE += [dict(w, part="A", poziom="INFO STP", opis="bryła A"),
                            dict(w, part="B", poziom="INFO STP", opis="bryła B"),
                            dict(w, part="B", poziom="BŁĄD", opis="błąd B")]
    sciezka = os.path.join(TMP, "raport_stp.xlsx")
    S.zapisz_raport_xlsx(sciezka)
    wb = openpyxl.load_workbook(sciezka)
    glowna = [r[2] for r in wb["Błędy i uwagi"].iter_rows(min_row=2, values_only=True) if r and r[4] == "INFO STP"]
    stp = [r[2] for r in wb["Model STP (info)"].iter_rows(min_row=2, values_only=True)]
    assert glowna == ["B"] and stp == ["A"], (glowna, stp)
    S.WSZYSTKIE_WIERSZE.clear()


def test_caly_folder():
    """Cały przebieg sprawdz() na folderze wydania: BOM + DXF + PDF z widokiem wektorowym części."""
    for nazwa, faza, oczekiwany, status in (("GE9001", 10, 0, "ZGODNY 1:1"), ("GE9002", 5, 1, "RÓŻNICE")):
        folder = os.path.join(TMP, nazwa)
        os.makedirs(folder)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["POS", "PART NUMBER", "TITLE", "", "BOM STRUCTURE", "QTY", "", "DESCRIPTION", "", "",
                   "MATERIAL", "M", "O", "P", "U", "C"])
        ws.append([1, "PG1", "SIDE PLATE", "", "Normal", 1, "", "PL 10 x 200 x 100", "", "", "S355"])
        wb.save(os.path.join(folder, f"{nazwa}.xlsx"))
        zapisz_dxf(f"{nazwa}__PG1__10mm__S355__1.dxf", prostokat_z_fazami(faza), folder=folder)
        # rysunek: widok 1:2 z fazami 10x10, notka przy widoku
        doc = fitz.open()
        strona = doc.new_page(width=842, height=595)
        sk = 72 / 25.4 / 2
        P = lambda x, y: fitz.Point(150 + x * sk, 400 - y * sk)  # noqa: E731
        for l in prostokat_z_fazami(10):
            strona.draw_line(P(*l[:2]), P(*l[2:4]), width=0.54)
        strona.insert_text(P(205, 105), "4x 10 x 45°", fontsize=8)
        strona.insert_text((150, 100), "VIEW1 ( 1 : 2 )", fontsize=9)
        strona.insert_text((600, 60), "PL 10 x 200 x 100", fontsize=9)
        strona.insert_text((600, 75), "S355", fontsize=9)
        doc.save(os.path.join(folder, "PG1__10mm__S355__1.pdf"))
        S.WSZYSTKIE_WIERSZE.clear()
        S.WERYFIKACJA_DXF.clear()
        assert S.sprawdz(folder) == oczekiwany, S.WSZYSTKIE_WIERSZE
        assert S.WERYFIKACJA_DXF[0]["status"] == status, S.WERYFIKACJA_DXF
        if oczekiwany:
            w = [w for w in S.WSZYSTKIE_WIERSZE if w["poziom"] == "BŁĄD"]
            assert w and w[0]["plik"] == "DXF" and w[0]["rodzaj"] in (
                "fazy DXF vs rysunek", "DXF vs rysunek 1:1 (nakładka)"), w
    S.zapisz_raport_xlsx(os.path.join(TMP, "raport.xlsx"))


if __name__ == "__main__":
    bledy = 0
    for nazwa, fn in list(globals().items()):
        if nazwa.startswith("test_"):
            try:
                fn()
                print("OK  ", nazwa)
            except Exception as e:
                bledy += 1
                print("FAIL", nazwa, "->", type(e).__name__, e)
    sys.exit(1 if bledy else 0)
